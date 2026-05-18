#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from common import (
    ABSTRACT_HEADER,
    CLASSIFICATION_HEADERS,
    EVIDENCE_PATH_HEADER,
    FIELD_MANUAL_STATUS_CONFIRMED,
    LOCAL_FILE_HEADER,
    PAPER_LINK_HEADER,
    PAPER_TITLE_HEADER,
    STATUS_BROWSER_PENDING,
    SUMMARY_HEADERS,
    WARNING_HEADER,
    WRITING_EVIDENCE_HEADER,
    WRITING_SECTION_HEADER,
    build_artifact_root,
    emit_json,
    ensure_dir,
    evidence_sufficiency_for_status,
    fieldguide_mapping_path_for_workbook,
    load_workflow_state,
    manifest_path_for_workbook,
    normalize_text,
    parse_json_file,
    project_file_paths,
    read_json_if_exists,
    relative_path,
    save_workflow_state,
    slugify,
    status_allows_full_backfill,
    write_json,
)


SCRIPT_DIR = Path(__file__).resolve().parent


def run_script(script_name: str, args: list[str], cwd: Path) -> dict[str, object]:
    cmd = [sys.executable, str(SCRIPT_DIR / script_name), *args]
    result = subprocess.run(
        cmd,
        cwd=str(cwd),
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or f"{script_name} failed")
    return json.loads(result.stdout)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def cell_value(ws, headers: dict[str, int], row: int, header: str) -> str:
    col = headers.get(header)
    if not col:
        return ""
    return normalize_text(ws.cell(row, col).value)


def parse_outline_sections(outline_text: str) -> list[dict[str, str]]:
    sections: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for raw_line in outline_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            if current:
                sections.append(current)
            heading = line[3:].strip()
            parts = heading.split(". ", 1)
            current = {
                "id": parts[0] if len(parts) > 1 else heading,
                "title": parts[1] if len(parts) > 1 else heading,
            }
    if current:
        sections.append(current)
    return sections


def choose_outline_section(title: str, values: dict[str, str], sections: list[dict[str, str]]) -> str:
    if values.get(WRITING_SECTION_HEADER):
        return values[WRITING_SECTION_HEADER]
    if values.get("功能映射层") == "Data":
        preferred = "5"
    elif values.get("核验后主要局限"):
        preferred = "6"
    elif values.get("范式映射") or values.get("协同粒度"):
        preferred = "3"
    else:
        preferred = "4"
    for section in sections:
        if section["id"] == preferred:
            return f"{section['id']}. {section['title']}"
    return f"{sections[0]['id']}. {sections[0]['title']}" if sections else "Pending outline mapping"


def choose_classification_value(field_name: str, existing_values: dict[str, str], title: str, abstract: str) -> str:
    existing = normalize_text(existing_values.get(field_name))
    if existing:
        return existing
    lowered = f"{title} {abstract}".lower()
    if field_name == "范式映射":
        if any(token in lowered for token in ("llm", "planner", "language-orchestrated", "orchestration", "code generation")):
            return "Paradigm II"
        if any(token in lowered for token in ("vla", "diffusion", "policy", "transformer", "flow matching")):
            return "Paradigm I"
    if field_name == "范式小类":
        if "diffusion" in lowered or "flow matching" in lowered:
            return "Diffusion-centered policies"
        if "code" in lowered or "skill" in lowered:
            return "Skill or code generation pipelines"
        if "planner" in lowered or "planning" in lowered:
            return "Planning-enhanced pipelines"
        if "vla" in lowered:
            return "End-to-end VLAs"
    if field_name == "功能映射层":
        if "dataset" in lowered or "benchmark" in lowered or "data" in lowered:
            return "Data"
        if "perception" in lowered or "vision" in lowered or "tactile" in lowered:
            return "Perception"
        if "plan" in lowered or "planner" in lowered:
            return "Planning"
        if "policy" in lowered or "action" in lowered or "control" in lowered:
            return "Policy-Action"
    if field_name == "协同粒度":
        if "coordination" in lowered or "dual-arm" in lowered or "bimanual" in lowered:
            return "Bimanual-coordination-level"
        if "task decomposition" in lowered or "instruction" in lowered:
            return "Task-level"
        if "trajectory" in lowered or "action" in lowered:
            return "Action-level"
    if field_name == "是否为双臂针对设计":
        if any(token in lowered for token in ("bimanual", "dual-arm", "two-arm")):
            return "native"
    return existing


def summarize_core_method(title: str, abstract: str) -> str:
    if not abstract:
        return ""
    sentences = [part.strip() for part in abstract.replace("\n", " ").split(".") if part.strip()]
    if not sentences:
        return abstract[:280].strip()
    return ". ".join(sentences[:2]).strip()[:360]


def summarize_tasks(abstract: str) -> str:
    if not abstract:
        return ""
    markers = ["dataset", "benchmark", "task", "real-world", "simulation", "robot"]
    pieces = [segment.strip() for segment in abstract.replace("\n", " ").split(".") if segment.strip()]
    selected = [piece for piece in pieces if any(marker in piece.lower() for marker in markers)]
    if not selected:
        return ""
    return ". ".join(selected[:2])[:320]


def summarize_results(abstract: str) -> str:
    if not abstract:
        return ""
    pieces = [segment.strip() for segment in abstract.replace("\n", " ").split(".") if segment.strip()]
    selected = [piece for piece in pieces if any(token in piece.lower() for token in ("improv", "%", "success", "outperform", "gain", "better"))]
    if not selected:
        return ""
    return ". ".join(selected[:2])[:320]


def summarize_limitations(title: str, abstract: str) -> str:
    if not abstract:
        return ""
    lowered = f"{title} {abstract}".lower()
    if "sim-to-real" in lowered:
        return "Still depends on simulator or transfer fidelity, so real-world coordination robustness may remain bounded."
    if "dataset" in lowered or "data" in lowered:
        return "Scalability gains still depend on data quality, coverage, and evidence that transfer survives outside the reported setting."
    return "Requires conservative follow-up with fulltext evidence to confirm scope boundaries, failure modes, and generalization limits."


def build_writing_argument(title: str, values: dict[str, str], abstract: str, evidence_sufficiency: str, status: str) -> str:
    if status in {"abstract_only", "unresolved", "mismatch_or_unverifiable", STATUS_BROWSER_PENDING}:
        return ""
    if evidence_sufficiency not in {"strong", "moderate"}:
        return ""
    method = values.get("核验后核心方法") or summarize_core_method(title, abstract)
    result = values.get("核验后关键结果") or summarize_results(abstract)
    layer = values.get("功能映射层") or "the manipulation stack"
    paradigm = values.get("范式映射") or "the current taxonomy"
    if not method:
        return ""
    sentence = f"{title} is a useful {paradigm} example at the {layer} layer because {method}"
    if result:
        sentence += f" It reports that {result}"
    sentence += " This makes it suitable as a writing-ready anchor once the cited evidence segments are confirmed."
    return sentence[:600]


def build_field_note(field_name: str, value: str, status: str, evidence_sufficiency: str, sections: list[dict[str, str]]) -> dict[str, str]:
    section_label = choose_outline_section(field_name, {WRITING_SECTION_HEADER: ""}, sections)
    note = {
        "value": value,
        "segments": "S1, S2",
        "decision_question": f"How should `{field_name}` be assigned under the confirmed survey taxonomy?",
        "derivation": "Derived from the paper's available fulltext-first evidence, with conservative fallback when evidence is weaker.",
        "causal_reasoning": "Describe the paper's inputs, key modules, outputs, and coordination mechanism, then explain why that evidence supports this label.",
        "exclusion_reason": "Record the nearest plausible alternative and why the paper does not cleanly satisfy that neighboring category.",
        "evidence_sufficiency": evidence_sufficiency,
        "writing_section": section_label,
        "writing_argument": "",
    }
    if status in {"abstract_only", "unresolved", "mismatch_or_unverifiable", STATUS_BROWSER_PENDING}:
        note["derivation"] = "Only weak or pending evidence is currently available, so classification is conservative and may preserve prior values."
        note["exclusion_reason"] = "Neighboring categories remain plausible until stronger evidence is collected."
    return note


def infer_field_notes(values: dict[str, str], status: str, sections: list[dict[str, str]], abstract: str, title: str) -> dict[str, dict[str, str]] | None:
    evidence_sufficiency = evidence_sufficiency_for_status(status)
    notes: dict[str, dict[str, str]] = {}
    for key in CLASSIFICATION_HEADERS:
        value = normalize_text(values.get(key))
        if value:
            notes[key] = build_field_note(key, value, status, evidence_sufficiency, sections)
    if values.get(WRITING_SECTION_HEADER):
        notes[WRITING_SECTION_HEADER] = build_field_note(WRITING_SECTION_HEADER, values[WRITING_SECTION_HEADER], status, evidence_sufficiency, sections)
    writing_argument = values.get(WRITING_EVIDENCE_HEADER)
    if writing_argument:
        note = build_field_note(WRITING_EVIDENCE_HEADER, writing_argument, status, evidence_sufficiency, sections)
        note["writing_argument"] = writing_argument
        notes[WRITING_EVIDENCE_HEADER] = note
    if values.get("核验后关键结果"):
        result_note = build_field_note("核验后关键结果", values["核验后关键结果"], status, evidence_sufficiency, sections)
        result_note["writing_argument"] = build_writing_argument(title, values, abstract, evidence_sufficiency, status)
        notes["核验后关键结果"] = result_note
    return notes or None


def write_temp_json(path: Path, payload: object) -> Path:
    write_json(path, payload)
    return path


def placeholder_source_payload(row: int, title: str, row_dir: Path, warning: str) -> dict[str, object]:
    source_path = row_dir / "source.md"
    if not source_path.exists():
        source_path.write_text(
            "\n".join(
                [
                    f"# {title}",
                    "",
                    "- Source label: placeholder",
                    f"- Warning: {warning}",
                    "",
                    "## Abstract",
                    "",
                    "No source artifact was available during this rebuild pass.",
                    "",
                ]
            )
            + "\n",
            encoding="utf-8",
        )
    return {
        "row": row,
        "title": title,
        "status": "unresolved",
        "resolved_url": None,
        "secondary_urls": [],
        "source_kind": "unresolved",
        "evidence_level": "unresolved",
        "allow_full_backfill": False,
        "warning": warning,
        "local_source": str(source_path),
    }


def finalize_browser_pending(workbook_dir: Path, row_dir: Path, source_json_path: Path) -> dict[str, object]:
    return run_script(
        "finalize_browser_capture.py",
        [
            "--source-json",
            str(source_json_path),
            "--row-asset-dir",
            str(row_dir),
        ],
        workbook_dir,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet-name", required=True)
    parser.add_argument("--row-start", type=int, default=2)
    parser.add_argument("--row-end", type=int)
    parser.add_argument("--overwrite-existing-evidence", action="store_true")
    parser.add_argument("--survey-mode", action="store_true")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    workbook_dir = workbook_path.parent
    artifact_root = ensure_dir(build_artifact_root(workbook_path))
    papers_dir = ensure_dir(artifact_root / "papers")
    evidence_dir = ensure_dir(artifact_root / "evidence")
    tmp_dir = ensure_dir(artifact_root / "snapshots" / "tmp")
    manifest_path = manifest_path_for_workbook(workbook_path)
    workflow_state = load_workflow_state(workbook_path)
    if args.survey_mode and normalize_text(workflow_state.get("field_manual_status")) != FIELD_MANUAL_STATUS_CONFIRMED:
        raise RuntimeError("Survey-mode batch rebuild requires a confirmed field manual.")

    outline_text = project_file_paths(workbook_path)["outline"].read_text(encoding="utf-8", errors="replace") if project_file_paths(workbook_path)["outline"].exists() else ""
    sections = parse_outline_sections(outline_text)

    wb = load_workbook(workbook_path)
    ws = wb[args.sheet_name]
    headers = header_map(ws)

    manifest = read_json_if_exists(manifest_path, [])
    if not isinstance(manifest, list):
        manifest = []
    manifest_by_row = {int(item["row"]): item for item in manifest if isinstance(item, dict) and "row" in item}

    updates: list[dict[str, object]] = []
    processed_manifest: list[dict[str, object]] = []
    browser_pending_rows: list[int] = []
    row_end = args.row_end or ws.max_row

    for row in range(args.row_start, row_end + 1):
        title = cell_value(ws, headers, row, PAPER_TITLE_HEADER)
        if not title:
            continue

        existing_evidence = cell_value(ws, headers, row, EVIDENCE_PATH_HEADER)
        if existing_evidence and not args.overwrite_existing_evidence:
            continue

        link = cell_value(ws, headers, row, PAPER_LINK_HEADER) or None
        abstract = cell_value(ws, headers, row, ABSTRACT_HEADER)
        row_dir = ensure_dir(papers_dir / f"{row:03d}-{slugify(title)}")

        try:
            source_payload = run_script(
                "fetch_paper_sources.py",
                ["--title", title, "--row", str(row), "--artifacts-dir", str(papers_dir), *(["--link", link] if link else [])],
                workbook_dir,
            )
        except Exception as exc:
            source_payload = placeholder_source_payload(row, title, row_dir, str(exc))

        source_json_path = write_temp_json(tmp_dir / f"{row:03d}-{slugify(title)}-source.json", source_payload)

        if normalize_text(source_payload.get("status")) == STATUS_BROWSER_PENDING:
            try:
                source_payload = finalize_browser_pending(workbook_dir, row_dir, source_json_path)
                source_json_path = write_temp_json(tmp_dir / f"{row:03d}-{slugify(title)}-source.json", source_payload)
            except Exception:
                browser_pending_rows.append(row)
                processed_manifest.append(
                    {
                        "row": row,
                        "title": title,
                        "local_source": "",
                        "evidence_path": "",
                        "status": source_payload.get("status"),
                        "warning": source_payload.get("warning"),
                        "allow_full_backfill": bool(source_payload.get("allow_full_backfill")),
                        "resolved_url": source_payload.get("resolved_url"),
                    }
                )
                continue

        if not source_payload.get("local_source"):
            source_payload = placeholder_source_payload(
                row,
                title,
                row_dir,
                normalize_text(source_payload.get("warning")) or "No local source artifact was produced.",
            )
            source_json_path = write_temp_json(tmp_dir / f"{row:03d}-{slugify(title)}-source.json", source_payload)

        source_rel = relative_path(Path(str(source_payload["local_source"])).resolve(), workbook_dir)
        status = normalize_text(source_payload.get("status"))
        evidence_sufficiency = evidence_sufficiency_for_status(status)

        existing_values = {header: cell_value(ws, headers, row, header) for header in CLASSIFICATION_HEADERS + SUMMARY_HEADERS + [WRITING_SECTION_HEADER, WRITING_EVIDENCE_HEADER]}
        values: dict[str, str] = {LOCAL_FILE_HEADER: source_rel}
        links: dict[str, str] = {LOCAL_FILE_HEADER: source_rel}

        warning_value = normalize_text(source_payload.get("warning"))
        if warning_value:
            values[WARNING_HEADER] = warning_value

        for header in CLASSIFICATION_HEADERS:
            chosen = choose_classification_value(header, existing_values, title, abstract)
            if chosen and (status_allows_full_backfill(status) or existing_values.get(header)):
                values[header] = chosen

        core_method = existing_values.get("核验后核心方法") or summarize_core_method(title, abstract)
        tasks = existing_values.get("核验后任务/数据集") or summarize_tasks(abstract)
        results = existing_values.get("核验后关键结果") or summarize_results(abstract)
        limitations = existing_values.get("核验后主要局限") or summarize_limitations(title, abstract)

        if status_allows_full_backfill(status):
            if core_method:
                values["核验后核心方法"] = core_method
            if tasks:
                values["核验后任务/数据集"] = tasks
            if results:
                values["核验后关键结果"] = results
            if limitations:
                values["核验后主要局限"] = limitations

        if args.survey_mode:
            writing_section = choose_outline_section(title, values, sections)
            writing_argument = build_writing_argument(title, values, abstract, evidence_sufficiency, status)
            if writing_section:
                values[WRITING_SECTION_HEADER] = writing_section
            if writing_argument:
                values[WRITING_EVIDENCE_HEADER] = writing_argument
            elif evidence_sufficiency in {"weak", "pending"}:
                values[WARNING_HEADER] = " | ".join(filter(None, [values.get(WARNING_HEADER, ""), "写作论据未自动生成，因证据仍偏弱或待确认"]))

        row_updates_path = write_temp_json(tmp_dir / f"{row:03d}-{slugify(title)}-updates.json", values)
        field_notes = infer_field_notes(values, status, sections, abstract, title)
        field_notes_path = None
        if field_notes:
            field_notes_path = write_temp_json(tmp_dir / f"{row:03d}-{slugify(title)}-field-notes.json", field_notes)

        evidence_payload = run_script(
            "build_evidence_files.py",
            [
                "--row",
                str(row),
                "--title",
                title,
                *(["--link", link] if link else []),
                "--source-json",
                str(source_json_path),
                "--evidence-dir",
                str(evidence_dir),
                "--workspace-root",
                str(workbook_dir),
                "--sheet-updates-json",
                str(row_updates_path),
                *(["--field-notes-json", str(field_notes_path)] if field_notes_path else []),
            ],
            workbook_dir,
        )

        evidence_rel = normalize_text(evidence_payload.get("evidence_path"))
        if evidence_rel:
            values[EVIDENCE_PATH_HEADER] = evidence_rel
            links[EVIDENCE_PATH_HEADER] = evidence_rel

        updates.append({"row": row, "values": values, "links": links})
        processed_manifest.append(
            {
                "row": row,
                "title": title,
                "local_source": source_rel,
                "evidence_path": evidence_rel,
                "status": source_payload.get("status"),
                "warning": source_payload.get("warning"),
                "allow_full_backfill": bool(source_payload.get("allow_full_backfill")),
                "resolved_url": source_payload.get("resolved_url"),
            }
        )

    updates_json_path = write_temp_json(tmp_dir / "workbook-updates.json", updates)
    workbook_update_args = [str(workbook_path), args.sheet_name, "--updates-json", str(updates_json_path)]
    if args.survey_mode:
        workbook_update_args.append("--survey-mode")
    workbook_update_payload = run_script("update_workbook.py", workbook_update_args, workbook_dir)

    for item in processed_manifest:
        manifest_by_row[int(item["row"])] = item
    final_manifest = [manifest_by_row[key] for key in sorted(manifest_by_row)]
    write_json(manifest_path, final_manifest)

    workflow_state["last_processed_rows"] = [item["row"] for item in processed_manifest]
    workflow_state["browser_pending_rows"] = browser_pending_rows
    workflow_state["last_batch_manifest_path"] = str(manifest_path)
    save_workflow_state(workbook_path, workflow_state)

    payload = {
        "workbook": str(workbook_path),
        "sheet_name": args.sheet_name,
        "processed_rows": [item["row"] for item in processed_manifest],
        "browser_pending_rows": browser_pending_rows,
        "manifest_path": str(manifest_path),
        "workbook_update": workbook_update_payload,
        "survey_mode": args.survey_mode,
    }
    emit_json(payload)


if __name__ == "__main__":
    main()
