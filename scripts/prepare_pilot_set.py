#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from common import (
    ABSTRACT_HEADER,
    CLASSIFICATION_HEADERS,
    EVIDENCE_PATH_HEADER,
    PAPER_LINK_HEADER,
    PAPER_TITLE_HEADER,
    PILOT_STATUS_READY,
    WRITING_SECTION_HEADER,
    emit_json,
    load_workflow_state,
    normalize_text,
    pilot_summary_path_for_workbook,
    save_workflow_state,
    write_json,
)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def cell_value(ws, headers: dict[str, int], row: int, header: str) -> str:
    col = headers.get(header)
    if not col:
        return ""
    return normalize_text(ws.cell(row, col).value)


def infer_reason(values: dict[str, str]) -> str:
    reasons: list[str] = []
    if values.get("范式映射"):
        reasons.append(f"covers paradigm {values['范式映射']}")
    if values.get("协同粒度"):
        reasons.append(f"targets {values['协同粒度']}")
    if values.get(EVIDENCE_PATH_HEADER):
        reasons.append("already has an evidence artifact")
    else:
        reasons.append("still needs evidence validation")
    if values.get(WRITING_SECTION_HEADER):
        reasons.append(f"already mapped to {values[WRITING_SECTION_HEADER]}")
    return ", ".join(reasons)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet-name", required=True)
    parser.add_argument("--row-start", type=int, default=2)
    parser.add_argument("--row-end", type=int)
    parser.add_argument("--max-samples", type=int, default=8)
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    wb = load_workbook(workbook_path, read_only=True)
    ws = wb[args.sheet_name]
    headers = header_map(ws)
    row_end = args.row_end or ws.max_row

    candidates: list[dict[str, object]] = []
    seen_keys: set[tuple[str, str]] = set()
    for row in range(args.row_start, row_end + 1):
        title = cell_value(ws, headers, row, PAPER_TITLE_HEADER)
        if not title:
            continue
        values = {
            "范式映射": cell_value(ws, headers, row, "范式映射"),
            "协同粒度": cell_value(ws, headers, row, "协同粒度"),
            EVIDENCE_PATH_HEADER: cell_value(ws, headers, row, EVIDENCE_PATH_HEADER),
            WRITING_SECTION_HEADER: cell_value(ws, headers, row, WRITING_SECTION_HEADER),
            PAPER_LINK_HEADER: cell_value(ws, headers, row, PAPER_LINK_HEADER),
            ABSTRACT_HEADER: cell_value(ws, headers, row, ABSTRACT_HEADER),
        }
        key = (values["范式映射"] or "unknown", values["协同粒度"] or "unknown")
        if key in seen_keys and len(candidates) >= args.max_samples:
            continue
        seen_keys.add(key)
        candidates.append(
            {
                "row": row,
                "title": title,
                "paradigm": values["范式映射"] or "unassigned",
                "coordination_granularity": values["协同粒度"] or "unassigned",
                "writing_section": values[WRITING_SECTION_HEADER] or "",
                "link": values[PAPER_LINK_HEADER] or "",
                "has_evidence": bool(values[EVIDENCE_PATH_HEADER]),
                "reason": infer_reason(values),
                "evidence_strength_target": "strong" if values[EVIDENCE_PATH_HEADER] else "pending",
            }
        )
        if len(candidates) >= args.max_samples:
            break

    summary = {
        "row_count": len(candidates),
        "rows": candidates,
        "coverage_notes": [
            "Pilot should cover at least two paradigm buckets when available.",
            "Pilot should include rows with and without existing evidence artifacts.",
            "Pilot should include at least one row whose taxonomy is still weak or ambiguous.",
        ],
    }
    write_json(pilot_summary_path_for_workbook(workbook_path), summary)

    state = load_workflow_state(workbook_path)
    state["pilot_status"] = PILOT_STATUS_READY
    state["pilot_summary_path"] = str(pilot_summary_path_for_workbook(workbook_path))
    state["pilot_sample_count"] = len(candidates)
    save_workflow_state(workbook_path, state)

    emit_json(
        {
            "workbook": str(workbook_path),
            "sheet_name": args.sheet_name,
            "pilot_summary_path": str(pilot_summary_path_for_workbook(workbook_path)),
            "sample_count": len(candidates),
            "rows": candidates,
        }
    )


if __name__ == "__main__":
    main()
