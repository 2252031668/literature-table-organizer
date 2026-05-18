#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from common import (
    FIELD_GUIDANCE_HEADER,
    FIELD_HEADER,
    FIELD_MANUAL_STATUS_NEEDS_CONFIRMATION,
    FIELD_VALUE_HEADER,
    fieldguide_mapping_path_for_workbook,
    load_workflow_state,
    markdown_has_unresolved_placeholders,
    normalize_text,
    parse_json_file,
    project_file_paths,
    save_workflow_state,
    write_json,
    emit_json,
)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def detect_legacy_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, str]]:
    wb = load_workbook(workbook_path, read_only=True)
    ws = wb[sheet_name]
    headers = header_map(ws)
    field_col = headers.get(FIELD_HEADER)
    guide_col = headers.get(FIELD_GUIDANCE_HEADER) or headers.get("说明")
    value_col = headers.get(FIELD_VALUE_HEADER) or headers.get("推荐取值+（例子说明）")
    rows: list[dict[str, str]] = []
    if not field_col:
        return rows
    for row_idx in range(2, ws.max_row + 1):
        field_name = normalize_text(ws.cell(row_idx, field_col).value)
        if not field_name:
            continue
        rows.append(
            {
                "field": field_name,
                "guidance": normalize_text(ws.cell(row_idx, guide_col).value) if guide_col else "",
                "recommended": normalize_text(ws.cell(row_idx, value_col).value) if value_col else "",
            }
        )
    return rows


def parse_outline_sections(outline_text: str) -> list[dict[str, str]]:
    sections: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for raw_line in outline_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            if current:
                sections.append(current)
            heading = line[3:].strip()
            parts = heading.split(". ", 1)
            current = {
                "id": parts[0] if len(parts) > 1 else heading,
                "title": parts[1] if len(parts) > 1 else heading,
                "writing_question": "",
                "evidence_needs": "",
            }
            continue
        if current and line.startswith("- Writing question:"):
            current["writing_question"] = line.split(":", 1)[1].strip()
        elif current and line.startswith("- Evidence needs:"):
            current["evidence_needs"] = line.split(":", 1)[1].strip()
    if current:
        sections.append(current)
    return sections


def choose_section(field_name: str, sections: list[dict[str, str]]) -> dict[str, str] | None:
    mapping = {
        "范式映射": "3",
        "范式小类": "3",
        "功能映射层": "3",
        "协同粒度": "3",
        "是否为双臂针对设计": "3",
        "核验后核心方法": "4",
        "核验后任务/数据集": "5",
        "核验后关键结果": "4",
        "核验后主要局限": "6",
        "写作引用章节": "4",
        "引用论据": "4",
    }
    preferred = mapping.get(field_name)
    if preferred:
        for item in sections:
            if item["id"] == preferred:
                return item
    return sections[0] if sections else None


def infer_field_role(field_name: str) -> str:
    if field_name in {"范式映射", "范式小类", "功能映射层", "协同粒度", "是否为双臂针对设计"}:
        return "Stabilize the survey taxonomy and boundary cases so that cross-paper grouping is reproducible."
    if field_name == "核验后核心方法":
        return "Summarize the system's actual input-module-output chain and coordination mechanism for writing."
    if field_name == "核验后任务/数据集":
        return "Anchor what the paper truly evaluated rather than what the abstract claims broadly."
    if field_name == "核验后关键结果":
        return "Capture quantitative or otherwise concrete evidence that supports comparison writing."
    if field_name == "核验后主要局限":
        return "Surface boundaries, bottlenecks, and future-work openings that matter to the survey argument."
    if field_name == "写作引用章节":
        return "Link the paper to the specific outline section it most directly supports."
    if field_name == "引用论据":
        return "Store a writing-ready argument that includes specific evidence, not a generic summary."
    return "Record a field that supports cross-paper survey writing and evidence-backed comparison."


def infer_decision_question(field_name: str) -> str:
    if field_name == "协同粒度":
        return "At what coordination level does the paper make its main contribution: task, action, or bimanual coordination?"
    if field_name == "功能映射层":
        return "Where does the paper's main foundation-model contribution enter the embodied stack?"
    if field_name == "是否为双臂针对设计":
        return "Is the system natively designed around bimanual coordination, or is bimanual use incidental?"
    if field_name == "写作引用章节":
        return "Which confirmed outline section most needs this paper as evidence or contrast?"
    if field_name == "引用论据":
        return "What writing-ready claim can this paper support once evidence sufficiency is checked?"
    return f"How should `{field_name}` be assigned under the confirmed survey taxonomy and why?"


def infer_positive_trigger(field_name: str) -> str:
    if field_name in {"范式映射", "范式小类", "功能映射层", "协同粒度"}:
        return "A stable label should be supported by explicit evidence about inputs, modules, outputs, or coordination logic."
    if field_name.startswith("核验后"):
        return "Use only evidence that appears in the paper's body, tables, figures, appendix, or trustworthy fulltext surrogate."
    return "Fill only when you can cite explicit evidence rather than paraphrasing the abstract."


def infer_negative_cases(field_name: str) -> str:
    if field_name == "协同粒度":
        return "Do not select a finer-grained coordination label when the paper only plans at task level or merely mentions coordination."
    if field_name == "功能映射层":
        return "Do not treat every use of a large model as policy/action if the paper mainly uses it for data, perception, or planning."
    if field_name == "引用论据":
        return "Avoid vague summary sentences that do not include concrete method differences, numeric results, or boundaries."
    return "Record the nearest plausible labels that were considered and explain why they were rejected."


def infer_evidence_requirement(field_name: str) -> str:
    if field_name == "引用论据":
        return "Prefer strong or moderate evidence with concrete numerical or mechanism-level support; weak evidence should trigger conservative wording."
    if field_name.startswith("核验后"):
        return "Prefer PDF/fulltext evidence that is traceable to a page, line, figure, or table region."
    return "Prefer fulltext evidence that reveals the paper's actual mechanism rather than abstract-level framing."


def infer_conservative_rule(field_name: str) -> str:
    if field_name == "引用论据":
        return "If evidence is weak or pending, leave the cell empty or write a conservative note plus warning instead of a confident argument."
    return "If evidence is insufficient, downgrade confidence, preserve old values when necessary, and record the warning explicitly."


def infer_writing_use(field_name: str, section: dict[str, str] | None) -> str:
    if not section:
        return "Pending confirmed outline alignment."
    return f"Supports section {section['id']}. {section['title']} by contributing evidence to: {section.get('writing_question') or 'the section argument.'}"


def build_manual(rows: list[dict[str, str]], outline_sections: list[dict[str, str]], gap_summary: dict[str, object]) -> str:
    lines = [
        "# Field Manual",
        "",
        "This manual is the authoritative decision protocol for survey-oriented classification and writing support.",
        "",
        "## Status",
        f"- Field manual status: {FIELD_MANUAL_STATUS_NEEDS_CONFIRMATION}",
        "",
        "## Gap summary",
        f"- Missing fields: {', '.join(gap_summary.get('missing_fields') or []) or 'None detected'}",
        f"- Ambiguous fields: {', '.join(gap_summary.get('ambiguous_fields') or []) or 'None detected'}",
        f"- Evidence-only dimensions: {', '.join(gap_summary.get('evidence_only_dimensions') or []) or 'None recorded'}",
        "",
        "## Decision entries",
        "",
    ]

    for row in rows:
        field_name = row["field"]
        section = choose_section(field_name, outline_sections)
        section_label = f"{section['id']}. {section['title']}" if section else "Pending outline alignment"
        lines.extend(
            [
                f"### {field_name}",
                f"- 字段: {field_name}",
                f"- 服务章节/写作问题: {section_label}",
                f"- 字段作用: {infer_field_role(field_name)}",
                f"- 判定问题: {infer_decision_question(field_name)}",
                f"- 推荐取值: {row['recommended'] or 'Use the confirmed taxonomy or section-aligned controlled vocabulary.'}",
                f"- 建议填写方式: {row['guidance'] or 'Prefer fulltext evidence first, then conservative secondary evidence if fulltext is unavailable.'}",
                f"- 正例/触发条件: {infer_positive_trigger(field_name)}",
                f"- 反例/易混淆项: {infer_negative_cases(field_name)}",
                f"- 关键证据要求: {infer_evidence_requirement(field_name)}",
                f"- 不足证据时如何处理: {infer_conservative_rule(field_name)}",
                "- 邻近类别排除规则: Record why the closest alternative label was not chosen, especially for taxonomy fields.",
                f"- 写作用途说明: {infer_writing_use(field_name, section)}",
                "",
            ]
        )
    return "\n".join(lines).strip() + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet-name", required=True)
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    rows = detect_legacy_rows(workbook_path, args.sheet_name)
    file_paths = project_file_paths(workbook_path)
    outline_text = file_paths["outline"].read_text(encoding="utf-8") if file_paths["outline"].exists() else ""
    gap_path = file_paths["field_gap_analysis"]
    gap_summary: dict[str, object] = {}
    if gap_path.exists():
        mapping_path = fieldguide_mapping_path_for_workbook(workbook_path)
        if mapping_path.exists():
            mapping_payload = parse_json_file(mapping_path)
            if isinstance(mapping_payload, dict):
                gap_summary = load_workflow_state(workbook_path).get("field_gap_summary", {})
        if not gap_summary:
            gap_summary = load_workflow_state(workbook_path).get("field_gap_summary", {})
    outline_sections = parse_outline_sections(outline_text)
    manual = build_manual(rows, outline_sections, gap_summary if isinstance(gap_summary, dict) else {})
    file_paths["field_manual"].parent.mkdir(parents=True, exist_ok=True)
    file_paths["field_manual"].write_text(manual, encoding="utf-8")

    state = load_workflow_state(workbook_path)
    state["fieldguide_sheet"] = args.sheet_name
    state["field_manual_status"] = FIELD_MANUAL_STATUS_NEEDS_CONFIRMATION
    state["field_manual_has_placeholders"] = markdown_has_unresolved_placeholders(manual)
    save_workflow_state(workbook_path, state)

    payload = {
        "workbook": str(workbook_path),
        "sheet_name": args.sheet_name,
        "legacy_entry_count": len(rows),
        "field_manual_path": str(file_paths["field_manual"]),
        "field_manual_status": FIELD_MANUAL_STATUS_NEEDS_CONFIRMATION,
        "requires_user_confirmation": True,
    }
    emit_json(payload)


if __name__ == "__main__":
    main()
