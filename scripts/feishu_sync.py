#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from common import col_to_letter, resolve_feishu_sheet_token, run_command


def get_remote_sheet_info(url: str) -> dict[str, object]:
    result = run_command(["lark-cli", "sheets", "+info", "--url", url])
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "Failed to query remote sheet info")
    return json.loads(result.stdout)


def unpack_remote_sheets(payload: dict[str, object]) -> list[dict[str, object]]:
    data = payload.get("data", {})
    sheets = data.get("sheets", {})
    if isinstance(sheets, dict):
        nested = sheets.get("sheets", [])
        if isinstance(nested, list):
            return [item for item in nested if isinstance(item, dict)]
    if isinstance(sheets, list):
        return [item for item in sheets if isinstance(item, dict)]
    top_level = payload.get("sheets", [])
    if isinstance(top_level, list):
        return [item for item in top_level if isinstance(item, dict)]
    return []


def workbook_used_range(ws) -> tuple[int, int]:
    max_row = ws.max_row
    max_col = ws.max_column
    while max_row > 1 and all(ws.cell(max_row, col).value in (None, "") for col in range(1, max_col + 1)):
        max_row -= 1
    while max_col > 1 and all(ws.cell(row, max_col).value in (None, "") for row in range(1, max_row + 1)):
        max_col -= 1
    return max_row, max_col


def sheet_matrix(ws, max_row: int, max_col: int) -> list[list[object]]:
    values: list[list[object]] = []
    for row_idx in range(1, max_row + 1):
        row_values = []
        for col_idx in range(1, max_col + 1):
            row_values.append(ws.cell(row_idx, col_idx).value)
        values.append(row_values)
    return values


def preview_sync(local_workbook: Path, url: str) -> dict[str, object]:
    remote = get_remote_sheet_info(url)
    remote_sheets = unpack_remote_sheets(remote)
    remote_by_title = {sheet.get("title"): sheet for sheet in remote_sheets if sheet.get("title")}

    wb = load_workbook(local_workbook, data_only=False)
    preview_items: list[dict[str, object]] = []
    for ws in wb.worksheets:
        max_row, max_col = workbook_used_range(ws)
        preview_items.append(
            {
                "sheet_name": ws.title,
                "remote_exists": ws.title in remote_by_title,
                "used_range": f"A1:{col_to_letter(max_col)}{max_row}",
                "remote_sheet_id": remote_by_title.get(ws.title, {}).get("sheet_id"),
            }
        )
    return {
        "spreadsheet_url": url,
        "local_workbook": str(local_workbook),
        "items": preview_items,
    }


def execute_sync(local_workbook: Path, url: str) -> dict[str, object]:
    remote = get_remote_sheet_info(url)
    remote_sheets = unpack_remote_sheets(remote)
    remote_by_title = {sheet.get("title"): sheet for sheet in remote_sheets if sheet.get("title")}

    wb = load_workbook(local_workbook, data_only=False)
    writes: list[dict[str, object]] = []
    spreadsheet_info = resolve_feishu_sheet_token(url)

    for ws in wb.worksheets:
        if ws.title not in remote_by_title:
            writes.append({"sheet_name": ws.title, "status": "skipped_missing_remote_sheet"})
            continue
        remote_sheet_id = remote_by_title[ws.title]["sheet_id"]
        max_row, max_col = workbook_used_range(ws)
        matrix = sheet_matrix(ws, max_row, max_col)
        range_ref = f"{remote_sheet_id}!A1:{col_to_letter(max_col)}{max_row}"
        cmd = [
            "lark-cli",
            "sheets",
            "+write",
            "--url",
            url,
            "--sheet-id",
            remote_sheet_id,
            "--range",
            range_ref,
            "--values",
            json.dumps(matrix, ensure_ascii=False),
        ]
        result = run_command(cmd, cwd=local_workbook.parent)
        if result.returncode != 0:
            writes.append(
                {
                    "sheet_name": ws.title,
                    "status": "error",
                    "message": result.stderr.strip() or result.stdout.strip(),
                }
            )
        else:
            writes.append(
                {
                    "sheet_name": ws.title,
                    "status": "written",
                    "range": range_ref,
                }
            )
    return {
        "spreadsheet_url": url,
        "spreadsheet_token": spreadsheet_info["token"],
        "local_workbook": str(local_workbook),
        "writes": writes,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--spreadsheet-url", required=True)
    parser.add_argument("--local-workbook", required=True)
    parser.add_argument("--mode", choices=["preview", "execute"], default="preview")
    args = parser.parse_args()

    local_workbook = Path(args.local_workbook).resolve()
    if args.mode == "preview":
        payload = preview_sync(local_workbook, args.spreadsheet_url)
    else:
        payload = execute_sync(local_workbook, args.spreadsheet_url)
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
