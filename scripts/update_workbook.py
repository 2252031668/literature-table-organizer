#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from common import (
    EVIDENCE_PATH_HEADER,
    LOCAL_FILE_HEADER,
    SYSTEM_HEADERS,
    WARNING_HEADER,
    WARNING_MISSING_TARGET,
    emit_json,
    normalize_text,
    parse_json_file,
)

HEADER_ALIASES = {
    "??????": LOCAL_FILE_HEADER,
    "?????": EVIDENCE_PATH_HEADER,
}


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            value = HEADER_ALIASES.get(value, value)
            mapping[value] = col_idx
    return mapping


def ensure_headers(ws) -> dict[str, int]:
    mapping = header_map(ws)
    next_col = ws.max_column + 1
    for header in SYSTEM_HEADERS:
        if header not in mapping:
            ws.cell(1, next_col).value = header
            mapping[header] = next_col
            next_col += 1
    return mapping


def ensure_header(mapping: dict[str, int], ws, header: str) -> int:
    if header in mapping:
        return mapping[header]
    index = ws.max_column + 1
    ws.cell(1, index).value = header
    mapping[header] = index
    return index


def set_local_hyperlink(ws, row: int, col: int, display_value: str, workbook_path: Path) -> bool:
    cell = ws.cell(row=row, column=col)
    cell.value = display_value
    if not display_value:
        cell.hyperlink = None
        return False
    target = (workbook_path.parent / display_value).resolve()
    if not target.exists():
        cell.hyperlink = None
        return False
    cell.hyperlink = display_value
    cell.style = "Hyperlink"
    return True


def apply_updates(ws, updates: list[dict[str, object]], workbook_path: Path) -> list[dict[str, object]]:
    mapping = ensure_headers(ws)
    warnings_added: list[dict[str, object]] = []
    for item in updates:
        row = int(item["row"])
        values: dict[str, object] = dict(item.get("values", {}))
        links: dict[str, object] = dict(item.get("links", {}))
        for header in set(values) | set(links):
            ensure_header(mapping, ws, header)

        row_warning_messages: list[str] = []
        for header, value in values.items():
            ws.cell(row=row, column=mapping[header]).value = value

        for header, value in links.items():
            link_value = normalize_text(value)
            link_written = set_local_hyperlink(ws, row, mapping[header], link_value, workbook_path)
            if not link_written and link_value:
                row_warning_messages.append(f"{WARNING_MISSING_TARGET}: {header}")

        if row_warning_messages:
            warning_col = ensure_header(mapping, ws, WARNING_HEADER)
            existing = normalize_text(ws.cell(row=row, column=warning_col).value)
            merged = " | ".join(filter(None, [existing, *row_warning_messages]))
            ws.cell(row=row, column=warning_col).value = merged
            warnings_added.append({"row": row, "warning": merged})
    return warnings_added


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", help="Workbook path")
    parser.add_argument("sheet_name", help="Primary worksheet name")
    parser.add_argument("--updates-json", help="JSON list of row updates")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    wb = load_workbook(workbook_path)
    ws = wb[args.sheet_name]
    ensure_headers(ws)

    updates = parse_json_file(Path(args.updates_json)) if args.updates_json else []
    warnings_added: list[dict[str, object]] = []
    if updates:
        warnings_added = apply_updates(ws, updates, workbook_path)

    wb.save(workbook_path)
    payload = {
        "workbook": str(workbook_path),
        "sheet_name": args.sheet_name,
        "headers": header_map(ws),
        "updated_rows": [item["row"] for item in updates],
        "warnings_added": warnings_added,
        "link_columns": [LOCAL_FILE_HEADER, EVIDENCE_PATH_HEADER],
    }
    emit_json(payload)


if __name__ == "__main__":
    main()
