#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from common import (
    FIELD_GUIDANCE_HEADER,
    FIELD_HEADER,
    FIELD_MANUAL_STATUS_DRAFT,
    FIELD_VALUE_HEADER,
    PILOT_STATUS_DRAFT,
    PROJECT_MODE_SURVEY,
    WRITING_EVIDENCE_HEADER,
    WRITING_SECTION_HEADER,
    emit_json,
    extract_markdown_headings,
    fieldguide_mapping_path_for_workbook,
    load_workflow_state,
    normalize_text,
    project_file_paths,
    read_draft_source,
    save_workflow_state,
    write_json,
)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def collect_fieldguide_entries(workbook_path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    if not sheet_name:
        return []
    wb = load_workbook(workbook_path, read_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    headers = header_map(ws)
    field_col = headers.get(FIELD_HEADER)
    guidance_col = headers.get(FIELD_GUIDANCE_HEADER) or headers.get("说明")
    value_col = headers.get(FIELD_VALUE_HEADER) or headers.get("推荐取值+（例子说明）")
    if not field_col:
        return []
    rows: list[dict[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        field_name = normalize_text(ws.cell(row_idx, field_col).value)
        if not field_name:
            continue
        rows.append(
            {
                "field": field_name,
                "guidance": normalize_text(ws.cell(row_idx, guidance_col).value) if guidance_col else "",
                "recommended": normalize_text(ws.cell(row_idx, value_col).value) if value_col else "",
            }
        )
    return rows


def classify_article_purpose(text: str) -> str:
    lowered = text.lower()
    if any(token in lowered for token in ("gap", "breakthrough", "positioning", "differentiate")):
        return "突破口论证"
    if "compare" in lowered or "comparison" in lowered:
        return "对比分析"
    if "survey" in lowered or "taxonomy" in lowered:
        return "分类综述"
    return "选题判断"


def infer_target_audience(topic: str, draft_text: str) -> str:
    lowered = f"{topic} {draft_text}".lower()
    if "survey" in lowered or "review" in lowered:
        return "Robotics and embodied AI researchers preparing a survey manuscript"
    return "Researchers exploring the topic landscape and evidence boundaries"


def infer_submission_positioning(topic: str, draft_text: str) -> str:
    lowered = f"{topic} {draft_text}".lower()
    if "comprehensive survey" in lowered or "survey" in lowered:
        return "Survey/review submission targeting robotics, embodied AI, or intelligent systems venues"
    return "Review-style manuscript or technical position paper"


def detect_existing_assets(field_rows: list[dict[str, str]], draft_text: str) -> dict[str, list[str]]:
    field_names = [row["field"] for row in field_rows]
    headings = extract_markdown_headings(draft_text)
    return {
        "field_names": field_names,
        "draft_headings": headings[:20],
    }


def infer_candidate_breakthroughs(topic: str, field_rows: list[dict[str, str]], draft_text: str) -> list[str]:
    lowered = f"{topic} {draft_text}".lower()
    candidates: list[str] = []
    if "bimanual" in lowered or "双臂" in lowered:
        candidates.append("Treat bimanual coordination as a first-class design variable rather than a larger version of general manipulation.")
    if "foundation" in lowered or "vla" in lowered or "llm" in lowered:
        candidates.append("Explain where foundation models enter the stack and how that changes coordination-aware system design.")
    if any("协同粒度" in row["field"] for row in field_rows):
        candidates.append("Use coordination granularity as a stronger organizing axis than model family names alone.")
    if not candidates:
        candidates.append("Position the survey around the gap between broad foundation-model reviews and coordination-specific manipulation demands.")
    return candidates


def build_outline_sections(topic: str, draft_headings: list[str], field_rows: list[dict[str, str]]) -> list[dict[str, object]]:
    field_names = [row["field"] for row in field_rows]
    sections = [
        {
            "id": "1",
            "title": "Introduction and Motivation",
            "writing_question": "Why does this topic require a dedicated survey now?",
            "evidence_needs": "Gap statements, representative progress, and why neighboring surveys are insufficient.",
            "fields": [],
        },
        {
            "id": "2",
            "title": "Scope and Positioning Against Existing Surveys",
            "writing_question": "What does this survey include, exclude, and differentiate itself from?",
            "evidence_needs": "Existing survey coverage, blind spots, and explicit boundary decisions.",
            "fields": [],
        },
        {
            "id": "3",
            "title": "Taxonomy and Decision Axes",
            "writing_question": "Which axes most meaningfully organize the literature for the target survey?",
            "evidence_needs": "Stable category definitions, boundary cases, and why neighboring labels are excluded.",
            "fields": [name for name in field_names if name in {"范式映射", "范式小类", "功能映射层", "协同粒度", "是否为双臂针对设计"}],
        },
        {
            "id": "4",
            "title": "Representative Methods and Coordination Mechanisms",
            "writing_question": "How do representative methods differ in inputs, modules, outputs, and coordination logic?",
            "evidence_needs": "Fulltext evidence for modules, coordination mechanism, numerical comparisons, and method boundaries.",
            "fields": [name for name in field_names if name in {"核验后核心方法", "核验后关键结果", WRITING_SECTION_HEADER, WRITING_EVIDENCE_HEADER}],
        },
        {
            "id": "5",
            "title": "Benchmarks, Datasets, and Evaluation Gaps",
            "writing_question": "How are systems evaluated, and what remains unstandardized?",
            "evidence_needs": "Datasets, tasks, metrics, real-world evidence, and open evaluation gaps.",
            "fields": [name for name in field_names if name in {"核验后任务/数据集", "核验后关键结果"}],
        },
        {
            "id": "6",
            "title": "Limitations and Outlook",
            "writing_question": "What bottlenecks remain, and which directions deserve emphasis?",
            "evidence_needs": "Failure modes, limitations, coordination bottlenecks, and research opportunities.",
            "fields": [name for name in field_names if name in {"核验后主要局限"}],
        },
    ]
    if draft_headings:
        sections[0]["draft_signals"] = draft_headings[:3]
        sections[1]["draft_signals"] = draft_headings[3:6]
    return sections


def build_project_brief(topic: str, workbook_path: Path, draft_info: dict[str, object], field_rows: list[dict[str, str]]) -> str:
    draft_text = str(draft_info.get("text") or "")
    audience = infer_target_audience(topic, draft_text)
    positioning = infer_submission_positioning(topic, draft_text)
    purpose = classify_article_purpose(draft_text)
    breakthroughs = infer_candidate_breakthroughs(topic, field_rows, draft_text)
    evidence_claims = [
        "Which taxonomy claims require fulltext evidence rather than abstract-only evidence.",
        "Which representative papers serve as counterexamples or boundary cases.",
        "Which quantitative comparisons directly support the survey's differentiating argument.",
    ]
    lines = [
        f"# Project Brief: {topic}",
        "",
        "## Project positioning",
        "",
        f"- Topic: {topic}",
        f"- Workbook: {workbook_path.name}",
        f"- Mode: {PROJECT_MODE_SURVEY}",
        f"- Target audience: {audience}",
        f"- Submission positioning: {positioning}",
        f"- Primary article purpose: {purpose}",
        f"- Workbook role: Build an evidence-backed survey taxonomy, representative-paper map, and writing-support table.",
        "",
        "## Current context",
        "",
    ]
    if draft_info.get("source"):
        lines.append(f"- Draft source: {draft_info['source']}")
        if draft_info.get("title"):
            lines.append(f"- Draft title: {draft_info['title']}")
    else:
        lines.append("- Draft source: None provided; the following analysis includes explicit inference where needed.")
    lines.extend(
        [
            f"- Legacy field-guide entries detected: {len(field_rows)}",
            "",
            "## Candidate breakthrough angles",
            "",
        ]
    )
    for index, item in enumerate(breakthroughs, start=1):
        lines.append(f"{index}. {item}")
    lines.extend(["", "## Claims that require strong evidence", ""])
    for claim in evidence_claims:
        lines.append(f"- {claim}")
    lines.extend(
        [
            "",
            "## Collaboration prompts",
            "",
            "- Which of the candidate breakthrough angles should become the survey's official main line?",
            "- Which boundary papers should be treated as calibration anchors rather than ordinary entries?",
            "- Which conclusions must be conservative unless we confirm PDF-level or fulltext-web evidence?",
            "",
        ]
    )
    return "\n".join(lines)


def build_related_survey_analysis(topic: str, draft_info: dict[str, object], field_rows: list[dict[str, str]]) -> str:
    draft_text = str(draft_info.get("text") or "")
    field_names = ", ".join(row["field"] for row in field_rows[:8]) or "No legacy field guide entries detected."
    candidate_breakthroughs = infer_candidate_breakthroughs(topic, field_rows, draft_text)
    lines = [
        f"# Related Survey Analysis: {topic}",
        "",
        "## Reading basis",
        "",
    ]
    if draft_info.get("source"):
        lines.append(f"- Survey draft or brief source: {draft_info['source']}")
    else:
        lines.append("- No draft source was provided; this analysis is inferred from the topic and workbook fields.")
    headings = draft_info.get("headings") or []
    if headings:
        lines.append(f"- Draft structural headings detected: {', '.join(headings[:8])}")
    lines.extend(
        [
            f"- Legacy field-guide signals: {field_names}",
            "",
            "## Existing-survey landscape to analyze",
            "",
            "- Broad foundation-model / robotics surveys: useful for stack-wide framing but often weak on coordination-specific taxonomy.",
            "- Foundation-model / manipulation surveys: useful for VLA, policy, and data-scaling framing but often not centered on dual-arm coordination.",
            "- Bimanual manipulation surveys: useful for coordination bottlenecks and task structure but often not organized around foundation-model entry points.",
            "",
            "## Gaps this survey should test explicitly",
            "",
            "1. Whether coordination should be treated as the main organizing axis rather than a subtopic under general manipulation.",
            "2. Whether the key distinction is where foundation models enter the embodied stack, not only what model family names are used.",
            "3. Whether representative boundary papers expose failure modes that neighboring surveys flatten away.",
            "",
            "## Candidate breakthrough ranking",
            "",
        ]
    )
    for index, item in enumerate(candidate_breakthroughs, start=1):
        lines.append(f"{index}. {item}")
    lines.extend(
        [
            "",
            "## Suggested main line",
            "",
            "Treat coordination-aware system design as the axis that differentiates this survey from both broad foundation-model surveys and general manipulation surveys.",
            "",
            "## Suggested counterexample line",
            "",
            "Highlight papers that challenge bigger-is-always-better or end-to-end-is-always-better assumptions by showing data-efficient, modular, or coordination-specialized alternatives.",
            "",
        ]
    )
    return "\n".join(lines)


def build_outline(topic: str, draft_info: dict[str, object], field_rows: list[dict[str, str]]) -> str:
    sections = build_outline_sections(topic, list(draft_info.get("headings") or []), field_rows)
    lines = [f"# Outline: {topic}", ""]
    for section in sections:
        lines.append(f"## {section['id']}. {section['title']}")
        lines.append(f"- Writing question: {section['writing_question']}")
        lines.append(f"- Evidence needs: {section['evidence_needs']}")
        fields = section.get("fields") or []
        if fields:
            lines.append(f"- Workbook fields primarily serving this section: {', '.join(fields)}")
        draft_signals = section.get("draft_signals") or []
        if draft_signals:
            lines.append(f"- Draft signals already present: {', '.join(draft_signals)}")
        lines.append("")
    lines.extend(
        [
            "## Cross-paper comparison rule",
            "- Sections 3-6 should prioritize cross-paper comparisons, not isolated abstract summaries.",
            "- Writing-support fields should link back to the section id and capture why a paper matters for that section.",
            "",
        ]
    )
    return "\n".join(lines)


def build_pilot_markdown(topic: str) -> str:
    return "\n".join(
        [
            f"# Pilot Calibration: {topic}",
            "",
            "## Goal",
            "- Calibrate 5-10 representative papers before any batch backfill.",
            "",
            "## Required record per paper",
            "- Initial classification",
            "- Why this classification was chosen",
            "- Neighbor category rejected and why",
            "- Evidence sufficiency",
            "- User correction",
            "- Rule learned",
            "- Field-manual clause updated",
            "- Outline section affected",
            "",
        ]
    )


def build_expansion_log_markdown(topic: str) -> str:
    return "\n".join(
        [
            f"# Paper Expansion Log: {topic}",
            "",
            "Record topic gaps, candidate papers, confirmation decisions, inserted rows, and follow-up evidence status.",
            "",
        ]
    )


def build_field_manual_stub(field_rows: list[dict[str, str]], sections: list[dict[str, object]]) -> str:
    section_lookup: list[str] = [f"{item['id']}. {item['title']}" for item in sections]
    lines = [
        "# Field Manual Draft",
        "",
        "This draft is the authoritative survey-oriented decision layer and must be confirmed before pilot or batch processing.",
        "",
        "## Status",
        f"- Field manual status: {FIELD_MANUAL_STATUS_DRAFT}",
        "",
        "## Candidate section map",
        "",
    ]
    for item in section_lookup:
        lines.append(f"- {item}")
    lines.extend(["", "## Draft entries", ""])
    for row in field_rows:
        field_name = row["field"]
        lines.extend(
            [
                f"### {field_name}",
                f"- 字段: {field_name}",
                "- 服务章节/写作问题: Pending alignment with confirmed outline section.",
                "- 字段作用: Capture a stable paper-level distinction that supports cross-paper writing rather than one-off table filling.",
                f"- 判定问题: Under the confirmed survey taxonomy, how should `{field_name}` be assigned and justified?",
                f"- 推荐取值: {row['recommended'] or 'Pending refinement from the confirmed survey taxonomy.'}",
                f"- 建议填写方式: {row['guidance'] or 'Use fulltext evidence first, then conservative secondary evidence if necessary.'}",
                "- 正例/触发条件: Add at least one explicit indicator from the full paper or trusted fulltext source.",
                "- 反例/易混淆项: Record neighboring labels that are tempting but unsupported by the paper's actual coordination mechanism.",
                "- 关键证据要求: Prefer PDF/fulltext evidence that reveals inputs, modules, outputs, and coordination logic.",
                "- 不足证据时如何处理: Downgrade to weak or pending, and avoid writing a strong survey-ready argument.",
                "- 邻近类别排除规则: Explain why other plausible labels are not selected.",
                "- 写作用途说明: Note which outline section this field supports and what kind of comparative sentence it should enable.",
                "",
            ]
        )
    return "\n".join(lines)


def build_field_gap_analysis(field_rows: list[dict[str, str]], sections: list[dict[str, object]]) -> tuple[str, dict[str, object]]:
    workbook_fields = [row["field"] for row in field_rows]
    outline_fields = []
    for section in sections:
        for field_name in section.get("fields") or []:
            if field_name not in outline_fields:
                outline_fields.append(field_name)
    missing = [field for field in [WRITING_SECTION_HEADER, WRITING_EVIDENCE_HEADER] + outline_fields if field and field not in workbook_fields]
    ambiguous = [row["field"] for row in field_rows if len(row["guidance"]) < 18 or len(row["recommended"]) < 18]
    redundant = [field for field in workbook_fields if field not in outline_fields and field not in {WRITING_SECTION_HEADER, WRITING_EVIDENCE_HEADER}]
    evidence_only = [
        "evidence_sufficiency",
        "exclusion_reason",
        "coordination_mechanism_anchor",
        "writing_argument_confidence",
    ]
    lines = [
        "# Field Gap Analysis",
        "",
        "## Missing fields or columns to ensure",
        "",
    ]
    for item in missing or ["None detected beyond optional writing-support refinements."]:
        lines.append(f"- {item}")
    lines.extend(["", "## Ambiguous legacy fields", ""])
    for item in ambiguous or ["None detected from the current heuristic."]:
        lines.append(f"- {item}")
    lines.extend(["", "## Redundant or low-priority workbook fields", ""])
    for item in redundant or ["None detected from the current heuristic."]:
        lines.append(f"- {item}")
    lines.extend(["", "## Evidence-only dimensions", ""])
    for item in evidence_only:
        lines.append(f"- {item}")
    return "\n".join(lines) + "\n", {
        "missing_fields": missing,
        "ambiguous_fields": ambiguous,
        "redundant_fields": redundant,
        "evidence_only_dimensions": evidence_only,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--topic", required=True)
    parser.add_argument("--draft-source")
    parser.add_argument("--fieldguide-sheet")
    parser.add_argument("--audience")
    parser.add_argument("--article-goal")
    parser.add_argument("--differentiation")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    draft_info = read_draft_source(args.draft_source)
    field_rows = collect_fieldguide_entries(workbook_path, args.fieldguide_sheet)
    sections = build_outline_sections(args.topic, list(draft_info.get("headings") or []), field_rows)
    file_paths = project_file_paths(workbook_path)

    brief = build_project_brief(args.topic, workbook_path, draft_info, field_rows)
    analysis = build_related_survey_analysis(args.topic, draft_info, field_rows)
    outline = build_outline(args.topic, draft_info, field_rows)
    pilot = build_pilot_markdown(args.topic)
    expansion = build_expansion_log_markdown(args.topic)
    manual = build_field_manual_stub(field_rows, sections)
    gap_markdown, gap_payload = build_field_gap_analysis(field_rows, sections)

    file_paths["brief"].parent.mkdir(parents=True, exist_ok=True)
    file_paths["brief"].write_text(brief + "\n", encoding="utf-8")
    file_paths["survey_analysis"].write_text(analysis + "\n", encoding="utf-8")
    file_paths["outline"].write_text(outline + "\n", encoding="utf-8")
    file_paths["pilot"].write_text(pilot + "\n", encoding="utf-8")
    file_paths["expansion_log"].write_text(expansion + "\n", encoding="utf-8")
    file_paths["field_manual"].write_text(manual + "\n", encoding="utf-8")
    file_paths["field_gap_analysis"].write_text(gap_markdown, encoding="utf-8")

    mapping_payload = {
        "fieldguide_sheet": args.fieldguide_sheet,
        "field_entries": field_rows,
        "outline_sections": [{"id": item["id"], "title": item["title"], "fields": item.get("fields") or []} for item in sections],
    }
    write_json(fieldguide_mapping_path_for_workbook(workbook_path), mapping_payload)

    state = load_workflow_state(workbook_path)
    state.update(
        {
            "mode": PROJECT_MODE_SURVEY,
            "topic": normalize_text(args.topic),
            "workbook": str(workbook_path),
            "draft_source": draft_info.get("source"),
            "draft_kind": draft_info.get("kind"),
            "draft_title": draft_info.get("title"),
            "fieldguide_sheet": args.fieldguide_sheet,
            "field_manual_status": FIELD_MANUAL_STATUS_DRAFT,
            "pilot_status": PILOT_STATUS_DRAFT,
            "project_files": {key: str(path) for key, path in file_paths.items()},
            "field_gap_summary": gap_payload,
        }
    )
    save_workflow_state(workbook_path, state)

    emit_json(
        {
            "mode": PROJECT_MODE_SURVEY,
            "workbook": str(workbook_path),
            "topic": normalize_text(args.topic),
            "draft_source": draft_info.get("source"),
            "fieldguide_sheet": args.fieldguide_sheet,
            "project_files": {key: str(path) for key, path in file_paths.items()},
            "field_gap_summary": gap_payload,
            "workflow_state_path": str((file_paths["brief"].parent / "workflow-state.json").resolve()),
        }
    )


if __name__ == "__main__":
    main()
