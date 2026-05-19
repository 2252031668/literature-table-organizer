#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from common import (
    PROJECT_MODE_SURVEY,
    STATUS_COMPLETE,
    STATUS_DOWNLOAD_FAILED,
    STATUS_NEEDS_LINK,
    STATUS_REVIEW_REQUIRED,
    emit_json,
    load_workflow_state,
    normalize_text,
    paper_reports_dir_for_workbook,
    papers_dir_for_workbook,
    project_file_paths,
    read_json_if_exists,
    row_analysis_dir_for_workbook,
    save_workflow_state,
    slugify,
    write_json,
    parse_row_spec,
)
from lib.link_search import find_paper_link
from lib.pdf_fetcher import download_pdf
from lib.project import build_framework, choose_primary_sheet, init_project
from lib.row_contracts import create_row_analysis_shell
from lib.workbook import append_rows, normalize_headers, writeback_rows, workbook_header_maps
from openpyxl import load_workbook


def command_init_project(args: argparse.Namespace) -> None:
    payload = init_project(Path(args.workbook), args.topic, args.draft_source, args.sheet_name)
    emit_json({"status": "ok", "command": "init-project", **payload})


def command_build_framework(args: argparse.Namespace) -> None:
    payload = build_framework(Path(args.workbook))
    emit_json({"status": "ok", "command": "build-framework", **payload})


def command_normalize_workbook_headers(args: argparse.Namespace) -> None:
    payload = normalize_headers(Path(args.workbook), Path(args.output).resolve() if args.output else None)
    emit_json({"status": "ok", "command": "normalize-workbook-headers", **payload})


def command_find_paper_link(args: argparse.Namespace) -> None:
    payload = find_paper_link(args.title)
    emit_json({"status": "ok", "command": "find-paper-link", **payload})


def command_pdf_download(args: argparse.Namespace) -> None:
    out_dir = Path(args.out).resolve()
    payload = download_pdf(args.url, out_dir, expected_title=args.title)
    emit_json({"command": "pdf-download", **payload})


def command_row_analyze(args: argparse.Namespace) -> None:
    payload = create_row_analysis_shell(Path(args.workbook), args.row, Path(args.pdf), args.sheet_name)
    emit_json(
        {
            "status": "ok",
            "command": "row-analyze",
            **payload,
            "message": "Read the PDF directly and fill the report in the workspace project directory. Do not edit the skill directory.",
        }
    )


def command_pilot_run(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook).resolve()
    rows = parse_row_spec(args.rows)
    review_queue_path = project_file_paths(workbook_path)["review_queue"]
    pilot_tasks_path = project_file_paths(workbook_path)["pilot_tasks"]

    wb = load_workbook(workbook_path, read_only=True)
    sheet_name = args.sheet_name or choose_primary_sheet(workbook_path)
    ws = wb[sheet_name]
    _, canonical_headers = workbook_header_maps(ws)
    papers_dir = papers_dir_for_workbook(workbook_path)

    tasks: list[dict[str, object]] = []
    markdown_lines = ["# Review Queue", "", "## Pilot Tasks", ""]
    for row in rows:
        def cell_for(key: str) -> str:
            col = canonical_headers.get(key)
            if not col:
                return ""
            return normalize_text(ws.cell(row, col).value)

        title = cell_for("paper_title")
        if not title:
            continue
        row_dir = papers_dir / f"{row:03d}-{slugify(title)}"
        pdf_path = row_dir / "paper.pdf"
        if pdf_path.exists():
            status = STATUS_REVIEW_REQUIRED
            blocker = ""
        elif cell_for("paper_link"):
            status = STATUS_DOWNLOAD_FAILED
            blocker = "PDF missing"
        else:
            status = STATUS_NEEDS_LINK
            blocker = "Link missing"
        task = {
            "row": row,
            "title": title,
            "pdf_path": str(pdf_path),
            "status": status,
            "blocker": blocker,
        }
        tasks.append(task)
        markdown_lines.append(f"- Row {row}: {title} | status={status} | blocker={blocker or 'none'}")

    review_queue_path.write_text("\n".join(markdown_lines).strip() + "\n", encoding="utf-8")
    write_json(pilot_tasks_path, {"sheet_name": sheet_name, "tasks": tasks})

    state = load_workflow_state(workbook_path)
    state["pilot_status"] = "ready"
    save_workflow_state(workbook_path, state)
    emit_json(
        {
            "status": "ok",
            "command": "pilot-run",
            "sheet_name": sheet_name,
            "task_count": len(tasks),
            "review_queue_path": str(review_queue_path),
            "pilot_tasks_path": str(pilot_tasks_path),
        }
    )


def command_batch_analyze(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook).resolve()
    wb = load_workbook(workbook_path, read_only=True)
    sheet_name = args.sheet_name or choose_primary_sheet(workbook_path)
    ws = wb[sheet_name]
    _, canonical_headers = workbook_header_maps(ws)
    row_start = args.row_start
    row_end = args.row_end or ws.max_row

    framework_path = project_file_paths(workbook_path)["framework"]
    review_queue_path = project_file_paths(workbook_path)["review_queue"]
    batch_tasks_path = project_file_paths(workbook_path)["batch_tasks"]
    papers_dir = papers_dir_for_workbook(workbook_path)
    reports_dir = paper_reports_dir_for_workbook(workbook_path)
    analysis_dir = row_analysis_dir_for_workbook(workbook_path)

    tasks: list[dict[str, object]] = []
    markdown_lines = ["# Review Queue", "", "## Batch Tasks", ""]
    for row in range(row_start, row_end + 1):
        def cell_for(key: str) -> str:
            col = canonical_headers.get(key)
            if not col:
                return ""
            return normalize_text(ws.cell(row, col).value)

        title = cell_for("paper_title")
        if not title:
            continue
        row_slug = slugify(title)
        pdf_path = papers_dir / f"{row:03d}-{row_slug}" / "paper.pdf"
        report_path = reports_dir / f"{row:03d}-{row_slug}.md"
        row_json_path = analysis_dir / f"{row:03d}-{row_slug}.json"

        blocker = ""
        if not framework_path.exists():
            status = STATUS_DOWNLOAD_FAILED
            blocker = "framework missing"
        elif not cell_for("paper_link") and not pdf_path.exists():
            status = STATUS_NEEDS_LINK
            blocker = "link and pdf missing"
        elif not pdf_path.exists():
            status = STATUS_DOWNLOAD_FAILED
            blocker = "pdf missing"
        elif row_json_path.exists():
            status = STATUS_COMPLETE
        else:
            status = STATUS_REVIEW_REQUIRED

        task = {
            "row": row,
            "title": title,
            "pdf_path": str(pdf_path),
            "report_path": str(report_path),
            "row_analysis_path": str(row_json_path),
            "status": status,
            "blocker": blocker,
        }
        tasks.append(task)
        markdown_lines.append(f"- Row {row}: {title} | status={status} | blocker={blocker or 'none'}")

        if args.prepare_shells and status == STATUS_REVIEW_REQUIRED and pdf_path.exists():
            create_row_analysis_shell(workbook_path, row, pdf_path, sheet_name)

    review_queue_path.write_text("\n".join(markdown_lines).strip() + "\n", encoding="utf-8")
    write_json(batch_tasks_path, {"sheet_name": sheet_name, "row_start": row_start, "row_end": row_end, "tasks": tasks})
    emit_json(
        {
            "status": "ok",
            "command": "batch-analyze",
            "sheet_name": sheet_name,
            "task_count": len(tasks),
            "review_queue_path": str(review_queue_path),
            "batch_tasks_path": str(batch_tasks_path),
        }
    )


def command_append_papers(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook).resolve()
    rows_json = Path(args.rows_json).resolve()
    payload = read_json_if_exists(rows_json, [])
    if not isinstance(payload, list):
        raise SystemExit("--rows-json must contain a JSON list")
    sheet_name = args.sheet_name or choose_primary_sheet(workbook_path)
    result = append_rows(workbook_path, sheet_name, payload)

    expansion_log_path = project_file_paths(workbook_path)["expansion_log"]
    lines = expansion_log_path.read_text(encoding="utf-8", errors="replace").splitlines() if expansion_log_path.exists() else ["# Expansion Log", ""]
    lines.append("## Append Run")
    lines.append("")
    lines.append("### Added")
    for item in result["added_rows"]:
        lines.append(f"- Row {item['row']}: {item['title']} | {item.get('link', '')}")
    if not result["added_rows"]:
        lines.append("- None")
    lines.append("")
    lines.append("### Skipped")
    for item in result.get("skipped_rows", []):
        lines.append(
            f"- Existing row {item['row']}: {item['title']} | {item.get('link', '')} | reason={item['reason']}"
            if item.get("row")
            else f"- No row: {item['title']} | {item.get('link', '')} | reason={item['reason']}"
        )
    if not result.get("skipped_rows"):
        lines.append("- None")
    expansion_log_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")

    emit_json({"status": "ok", "command": "append-papers", **result, "expansion_log_path": str(expansion_log_path)})


def command_writeback_xlsx(args: argparse.Namespace) -> None:
    payload = writeback_rows(Path(args.workbook), args.sheet_name, preserve_headers=args.preserve_headers and not args.normalize_to_english)
    emit_json({"status": "ok", "command": "writeback-xlsx", **payload})


def command_reset_project(args: argparse.Namespace) -> None:
    workbook_path = Path(args.workbook).resolve()
    reports_dir = paper_reports_dir_for_workbook(workbook_path)
    analysis_dir = row_analysis_dir_for_workbook(workbook_path)
    files = project_file_paths(workbook_path)
    removed: list[str] = []

    for path in sorted(reports_dir.glob("*")):
        path.unlink()
        removed.append(str(path))
    for path in sorted(analysis_dir.glob("*")):
        path.unlink()
        removed.append(str(path))
    for key in ["review_queue", "pilot_tasks", "batch_tasks"]:
        path = files[key]
        if path.exists():
            path.unlink()
            removed.append(str(path))

    emit_json({"status": "ok", "command": "reset-project", "removed": removed})


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Single-entry survey workflow CLI for literature-table-organizer")
    subparsers = parser.add_subparsers(dest="command", required=True)

    init_parser = subparsers.add_parser("init-project")
    init_parser.add_argument("--topic", required=True)
    init_parser.add_argument("--workbook", required=True)
    init_parser.add_argument("--draft-source")
    init_parser.add_argument("--sheet-name")
    init_parser.set_defaults(func=command_init_project)

    framework_parser = subparsers.add_parser("build-framework")
    framework_parser.add_argument("--workbook", required=True)
    framework_parser.set_defaults(func=command_build_framework)

    normalize_parser = subparsers.add_parser("normalize-workbook-headers")
    normalize_parser.add_argument("--workbook", required=True)
    normalize_parser.add_argument("--output")
    normalize_parser.set_defaults(func=command_normalize_workbook_headers)

    find_link_parser = subparsers.add_parser("find-paper-link")
    find_link_parser.add_argument("--title", required=True)
    find_link_parser.set_defaults(func=command_find_paper_link)

    pdf_parser = subparsers.add_parser("pdf-download")
    pdf_parser.add_argument("--url", required=True)
    pdf_parser.add_argument("--out", required=True)
    pdf_parser.add_argument("--title")
    pdf_parser.set_defaults(func=command_pdf_download)

    row_parser = subparsers.add_parser("row-analyze")
    row_parser.add_argument("--workbook", required=True)
    row_parser.add_argument("--row", type=int, required=True)
    row_parser.add_argument("--pdf", required=True)
    row_parser.add_argument("--sheet-name")
    row_parser.set_defaults(func=command_row_analyze)

    pilot_parser = subparsers.add_parser("pilot-run")
    pilot_parser.add_argument("--workbook", required=True)
    pilot_parser.add_argument("--rows", required=True)
    pilot_parser.add_argument("--sheet-name")
    pilot_parser.set_defaults(func=command_pilot_run)

    batch_parser = subparsers.add_parser("batch-analyze")
    batch_parser.add_argument("--workbook", required=True)
    batch_parser.add_argument("--row-start", type=int, required=True)
    batch_parser.add_argument("--row-end", type=int)
    batch_parser.add_argument("--sheet-name")
    batch_parser.add_argument("--prepare-shells", action="store_true")
    batch_parser.set_defaults(func=command_batch_analyze)

    append_parser = subparsers.add_parser("append-papers")
    append_parser.add_argument("--workbook", required=True)
    append_parser.add_argument("--rows-json", required=True)
    append_parser.add_argument("--sheet-name")
    append_parser.set_defaults(func=command_append_papers)

    writeback_parser = subparsers.add_parser("writeback-xlsx")
    writeback_parser.add_argument("--workbook", required=True)
    writeback_parser.add_argument("--sheet-name")
    writeback_parser.add_argument("--preserve-headers", action="store_true")
    writeback_parser.add_argument("--normalize-to-english", action="store_true")
    writeback_parser.set_defaults(func=command_writeback_xlsx)

    reset_parser = subparsers.add_parser("reset-project")
    reset_parser.add_argument("--workbook", required=True)
    reset_parser.set_defaults(func=command_reset_project)

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
