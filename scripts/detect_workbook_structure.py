#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from common import (
    FIELD_COL_KEYS,
    GUIDANCE_COL_KEYS,
    OPTIONAL_ABSTRACT_HEADERS,
    OPTIONAL_LINK_HEADERS,
    PRIMARY_TITLE_HEADERS,
    SYSTEM_HEADERS,
    VALUE_COL_KEYS,
    normalize_key,
    normalize_text,
)


def worksheet_headers(ws, header_row: int = 1, scan_cols: int = 80) -> list[str]:
    return [normalize_text(ws.cell(header_row, col_idx).value) for col_idx in range(1, scan_cols + 1)]


def score_primary_sheet(headers: list[str], ws) -> int:
    lowered = {normalize_key(h) for h in headers if h}
    has_title = bool(lowered & PRIMARY_TITLE_HEADERS)
    if not has_title:
        return 0
    score = 0
    score += 10
    if lowered & OPTIONAL_LINK_HEADERS:
        score += 3
    if lowered & OPTIONAL_ABSTRACT_HEADERS:
        score += 2
    non_empty_headers = len([header for header in headers if header])
    if non_empty_headers >= 3:
        score += 2
    if ws.max_row > 3:
        score += 1
    return score


def score_fieldguide_sheet(headers: list[str]) -> tuple[int, dict[str, int | None]]:
    lowered = {normalize_key(h) for h in headers if h}
    matched = {
        "field": 5 if lowered & FIELD_COL_KEYS else 0,
        "guidance": 5 if lowered & GUIDANCE_COL_KEYS else 0,
        "values": 5 if lowered & VALUE_COL_KEYS else 0,
    }
    return sum(matched.values()), {key: (1 if value else None) for key, value in matched.items()}


def detect_existing_columns(headers: list[str]) -> dict[str, int | None]:
    existing = {header: None for header in SYSTEM_HEADERS}
    for idx, header in enumerate(headers, start=1):
        if header in existing:
            existing[header] = idx
    return existing


def build_candidate_payload(ws, headers: list[str]) -> dict[str, object]:
    return {
        "sheet_name": ws.title,
        "headers": headers,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "existing_columns": detect_existing_columns(headers),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", help="Path to workbook")
    args = parser.parse_args()

    workbook = Path(args.workbook).resolve()
    wb = load_workbook(workbook, read_only=True, data_only=False)
    primary_candidates: list[dict[str, object]] = []
    fieldguide_candidates: list[dict[str, object]] = []

    for ws in wb.worksheets:
        headers = worksheet_headers(ws)
        primary_score = score_primary_sheet(headers, ws)
        fieldguide_score, fieldguide_semantics = score_fieldguide_sheet(headers)

        if primary_score > 0:
            item = build_candidate_payload(ws, headers)
            item["score"] = primary_score
            primary_candidates.append(item)

        if fieldguide_score > 0:
            item = build_candidate_payload(ws, headers)
            item["score"] = fieldguide_score
            item["semantic_hits"] = fieldguide_semantics
            fieldguide_candidates.append(item)

    primary_candidates.sort(key=lambda item: (int(item["score"]), int(item["max_row"])), reverse=True)
    fieldguide_candidates.sort(key=lambda item: (int(item["score"]), int(item["max_row"])), reverse=True)

    payload = {
        "workbook": str(workbook),
        "primary_candidates": primary_candidates[:3],
        "fieldguide_candidates": fieldguide_candidates[:3],
        "needs_user_choice": len(primary_candidates) > 1,
        "fieldguide_complete": bool(fieldguide_candidates and int(fieldguide_candidates[0]["score"]) >= 15),
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
