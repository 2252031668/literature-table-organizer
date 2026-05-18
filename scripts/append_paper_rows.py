#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from common import (
    ABSTRACT_HEADER,
    PAPER_LINK_HEADER,
    PAPER_TITLE_HEADER,
    WRITING_EVIDENCE_HEADER,
    WRITING_SECTION_HEADER,
    emit_json,
    normalize_text,
)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def ensure_header(ws, mapping: dict[str, int], header: str) -> int:
    if header in mapping:
        return mapping[header]
    col = ws.max_column + 1
    ws.cell(1, col).value = header
    mapping[header] = col
    return col


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet-name", required=True)
    parser.add_argument("--rows-json", required=True)
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    rows_path = Path(args.rows_json).resolve()
    payload = __import__("json").loads(rows_path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, list):
        raise SystemExit("--rows-json must contain a JSON list")

    wb = load_workbook(workbook_path)
    ws = wb[args.sheet_name]
    mapping = header_map(ws)
    ensure_header(ws, mapping, WRITING_SECTION_HEADER)
    ensure_header(ws, mapping, WRITING_EVIDENCE_HEADER)

    added_rows: list[dict[str, object]] = []
    for item in payload:
        title = normalize_text(item.get("title"))
        if not title:
            continue
        row_idx = ws.max_row + 1
        ws.cell(row_idx, mapping[PAPER_TITLE_HEADER]).value = title
        if PAPER_LINK_HEADER in mapping:
            ws.cell(row_idx, mapping[PAPER_LINK_HEADER]).value = normalize_text(item.get("link"))
        if ABSTRACT_HEADER in mapping:
            ws.cell(row_idx, mapping[ABSTRACT_HEADER]).value = normalize_text(item.get("abstract"))
        ws.cell(row_idx, mapping[WRITING_SECTION_HEADER]).value = normalize_text(item.get("writing_section"))
        ws.cell(row_idx, mapping[WRITING_EVIDENCE_HEADER]).value = normalize_text(item.get("writing_argument"))
        added_rows.append(
            {
                "row": row_idx,
                "title": title,
                "link": normalize_text(item.get("link")),
                "writing_section": normalize_text(item.get("writing_section")),
            }
        )

    wb.save(workbook_path)
    emit_json(
        {
            "workbook": str(workbook_path),
            "sheet_name": args.sheet_name,
            "added_rows": added_rows,
        }
    )


if __name__ == "__main__":
    main()
