#!/usr/bin/env python3
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import load_workbook

from common import (
    CLASSIFICATION_HEADERS,
    EVIDENCE_PATH_HEADER,
    LOCAL_FILE_HEADER,
    PAPER_LINK_HEADER,
    PAPER_TITLE_HEADER,
    SUMMARY_HEADERS,
    WARNING_HEADER,
    WRITING_EVIDENCE_HEADER,
    WRITING_SECTION_HEADER,
    build_artifact_root,
    emit_json,
    normalize_text,
)


def header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if value:
            mapping[value] = col_idx
    return mapping


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet-name", required=True)
    parser.add_argument("--row-start", type=int, default=2)
    parser.add_argument("--row-end", type=int)
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    artifact_root = build_artifact_root(workbook_path)
    evidence_dir = artifact_root / "evidence"
    if evidence_dir.exists():
        shutil.rmtree(evidence_dir)

    wb = load_workbook(workbook_path)
    ws = wb[args.sheet_name]
    headers = header_map(ws)
    row_end = args.row_end or ws.max_row
    cleared_rows: list[int] = []
    clear_headers = CLASSIFICATION_HEADERS + SUMMARY_HEADERS + [EVIDENCE_PATH_HEADER, WARNING_HEADER, WRITING_SECTION_HEADER, WRITING_EVIDENCE_HEADER]

    for row in range(args.row_start, row_end + 1):
        if not normalize_text(ws.cell(row, headers.get(PAPER_TITLE_HEADER, 0)).value if headers.get(PAPER_TITLE_HEADER) else ""):
            continue
        for header in clear_headers:
            col = headers.get(header)
            if col:
                ws.cell(row, col).value = None
                ws.cell(row, col).hyperlink = None
        local_col = headers.get(LOCAL_FILE_HEADER)
        if local_col:
            local_value = normalize_text(ws.cell(row, local_col).value)
            ws.cell(row, local_col).value = local_value
        cleared_rows.append(row)

    wb.save(workbook_path)
    emit_json(
        {
            "workbook": str(workbook_path),
            "sheet_name": args.sheet_name,
            "artifact_root": str(artifact_root),
            "removed_evidence_dir": str(evidence_dir),
            "cleared_rows": cleared_rows,
            "preserved_columns": [PAPER_TITLE_HEADER, PAPER_LINK_HEADER, "摘要", LOCAL_FILE_HEADER],
        }
    )


if __name__ == "__main__":
    main()
