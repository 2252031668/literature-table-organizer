from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from common import (
    FRAMEWORK_STATUS_DRAFT,
    PILOT_STATUS_DRAFT,
    PROJECT_MODE_SURVEY,
    canonical_key_for_header,
    ensure_dir,
    load_workflow_state,
    normalize_text,
    paper_reports_dir_for_workbook,
    papers_dir_for_workbook,
    project_dir_for_workbook,
    project_file_paths,
    read_draft_source,
    row_analysis_dir_for_workbook,
    save_workflow_state,
    write_json,
)


def choose_primary_sheet(workbook_path: Path) -> str:
    wb = load_workbook(workbook_path, read_only=True)
    best_name = wb.sheetnames[0]
    best_score = -1
    for name in wb.sheetnames:
        ws = wb[name]
        headers = [normalize_text(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
        score = sum(1 for header in headers if canonical_key_for_header(header))
        if score > best_score:
            best_score = score
            best_name = name
    return best_name


def default_framework_text(topic: str, workbook_path: Path, draft_info: dict[str, object], sheet_name: str) -> str:
    draft_summary = normalize_text(draft_info.get("title")) or "未提供草稿"
    draft_source = normalize_text(draft_info.get("source")) or "None"
    inferred_headings = draft_info.get("headings") or []
    heading_lines = "\n".join(f"- {item}" for item in inferred_headings[:6]) or "- 暂无可用草稿结构"
    return f"""# Framework

## Survey Context
- topic: {topic}
- workbook: {workbook_path.name}
- primary_sheet: {sheet_name}
- target_reader: 待讨论
- scope: 待讨论
- boundary: 待讨论
- differentiation: 待讨论
- draft_source: {draft_source}
- draft_summary: {draft_summary}

## Draft Signals
以下内容仅用于帮助快速继承已有草稿语境，需要继续人工整理：
{heading_lines}

## Compact Outline
### 1. Introduction and Motivation
- purpose: 说明为什么这个综述值得单独写，以及它相对已有综述的独特价值。
- evidence_need: 代表性进展、已有综述不足、领域转折点、任务需求变化。

### 2. Scope and Boundary
- purpose: 明确纳入范围、排除范围，以及哪些论文只作为边界/对照样本。
- evidence_need: 综述对比、方法边界、任务边界、foundation model 相关性。

### 3. Taxonomy and Classification Axes
- purpose: 定义工作簿中的分类字段，以及每个字段服务于哪类写作判断。
- evidence_need: 方法机制、输入输出接口、双臂角色关系、规划显式性。

### 4. Representative Systems and Comparative Analysis
- purpose: 组织代表论文，比较它们在哪一层真正改变了双臂问题。
- evidence_need: 系统框图、方法段落、实验设置、关键结果、失败案例。

### 5. Evaluation Protocols, Limitations, and Outlook
- purpose: 总结评测方式、局限、未解决问题和未来方向。
- evidence_need: 结果表、消融实验、局限讨论、作者展望。

## Classification Rules
### paradigm_mapping
- chinese_name: 范式映射
- allowed_values: 在项目讨论中补全，例如 `Paradigm I`, `Paradigm II`, `pending-boundary`
- minimum_evidence: 至少需要方法级证据，不允许只依据摘要或标题关键词。
- exclusion_rule: 必须说明为什么不是最接近的另一范式。
- fallback_rule: 证据不足时留空或标记 `pending-boundary`。

### paradigm_subtype
- chinese_name: 范式小类
- allowed_values: 按本项目最终 taxonomy 补全。
- minimum_evidence: 需要能说清论文的核心机制，而不是只看任务名。
- exclusion_rule: 需要排除邻近 subtype。
- fallback_rule: 若当前证据不足，则留空。

### functional_layer
- chinese_name: 功能映射层
- allowed_values: 建议最终只保留一个主层，回答“这篇论文最主要改变双臂问题的位置在哪里”。
- minimum_evidence: 需要明确其主要创新落点，例如 perception / data / policy / coordination interface。
- exclusion_rule: 必须解释为什么不是其他相邻层。
- fallback_rule: 若主要落点不清晰，则留空并警告。

### planning_granularity
- chinese_name: 规划粒度
- allowed_values: 仅在论文显式涉及 planner、task decomposition、role allocation 或 dual-arm coordination plan 时填写。
- minimum_evidence: 需要正文中明确描述规划模块或显式规划过程。
- exclusion_rule: 不能因为存在动作序列输出就默认视为规划。
- fallback_rule: 若无显式规划，保持为空。

### bimanual_design_type
- chinese_name: 是否为双臂针对设计
- allowed_values: 在项目讨论中补全，例如 `native`, `adapted`, `generalist`, `boundary`
- minimum_evidence: 需要系统设计或实验定位能支撑这个判断。
- exclusion_rule: 需要区分是否真的为双臂特化，而非一般操作模型直接迁移。
- fallback_rule: 证据不充分时保守填写或留空。

## Reading Report Template
### Part A - Chinese Reading Notes
- Basic Information
- 这篇论文解决什么问题
- 方法是什么，亮点是什么
- 实验怎么做
- 结果最关键的点是什么
- 局限和展望
- 这篇论文对本综述有什么意义

### Part B - English Field Decisions
- field
- decision
- reasoning
- confidence
- warning

## Workbook Mapping
- `row-analysis/*.json` is the only direct source for workbook writeback.
- `writing_section` must refer to a real section defined in this framework.
- `writing_argument` should be writing-ready, specific, and traceable to the PDF.
- Guarded fields should stay blank when the PDF does not provide enough support.
"""


def default_framework_schema() -> dict[str, object]:
    return {
        "report_structure": {
            "part_a_language": "zh-CN",
            "part_b_language": "en",
        },
        "required_sections": [
            "Survey Context",
            "Compact Outline",
            "Classification Rules",
            "Reading Report Template",
            "Workbook Mapping",
        ],
        "status_model": {
            "framework": [FRAMEWORK_STATUS_DRAFT, "confirmed"],
            "pilot": [PILOT_STATUS_DRAFT, "ready", "confirmed"],
        },
    }


def init_project(workbook_path: Path, topic: str, draft_source: str | None, sheet_name: str | None) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    chosen_sheet = sheet_name or choose_primary_sheet(workbook_path)
    draft_info = read_draft_source(draft_source)
    framework_text = default_framework_text(topic, workbook_path, draft_info, chosen_sheet)

    ensure_dir(project_dir_for_workbook(workbook_path))
    ensure_dir(papers_dir_for_workbook(workbook_path))
    ensure_dir(paper_reports_dir_for_workbook(workbook_path))
    ensure_dir(row_analysis_dir_for_workbook(workbook_path))

    files = project_file_paths(workbook_path)
    files["framework"].write_text(framework_text, encoding="utf-8")
    write_json(files["framework_schema"], default_framework_schema())
    if not files["project_brief"].exists():
        files["project_brief"].write_text("# Project Brief\n\n- Pending discussion.\n", encoding="utf-8")
    if not files["survey_gap_analysis"].exists():
        files["survey_gap_analysis"].write_text("# Survey Gap Analysis\n\n- Pending comparison against related surveys.\n", encoding="utf-8")
    if not files["pilot_calibration"].exists():
        files["pilot_calibration"].write_text("# Pilot Calibration\n\n- Pending pilot discussion.\n", encoding="utf-8")
    if not files["review_queue"].exists():
        files["review_queue"].write_text("# Review Queue\n\n- No pending items.\n", encoding="utf-8")
    if not files["expansion_log"].exists():
        files["expansion_log"].write_text("# Expansion Log\n\n- No paper additions recorded yet.\n", encoding="utf-8")

    state = load_workflow_state(workbook_path)
    state.update(
        {
            "mode": PROJECT_MODE_SURVEY,
            "framework_status": FRAMEWORK_STATUS_DRAFT,
            "pilot_status": PILOT_STATUS_DRAFT,
            "topic": normalize_text(topic),
            "workbook": str(workbook_path),
            "primary_sheet": chosen_sheet,
            "draft_source": draft_info.get("source"),
            "project_dir": str(project_dir_for_workbook(workbook_path)),
        }
    )
    save_workflow_state(workbook_path, state)

    return {
        "workbook": str(workbook_path),
        "primary_sheet": chosen_sheet,
        "project_files": {key: str(path) for key, path in files.items()},
    }


def build_framework(workbook_path: Path) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    files = project_file_paths(workbook_path)
    if not files["framework"].exists():
        raise FileNotFoundError("framework.md does not exist. Run init-project first.")
    state = load_workflow_state(workbook_path)
    state["framework_status"] = FRAMEWORK_STATUS_DRAFT
    save_workflow_state(workbook_path, state)
    return {
        "framework_path": str(files["framework"]),
        "framework_schema_path": str(files["framework_schema"]),
        "message": "Edit framework.md in the workspace project directory. Do not edit files inside the skill directory.",
    }
