from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from common import (
    CANONICAL_HEADER_ORDER,
    WARNING_HEADER,
    canonical_fieldguide_key,
    canonical_key_for_header,
    detect_header_language,
    normalize_text,
    preferred_header_for_key,
    project_file_paths,
    read_json_if_exists,
    relative_path,
    write_json,
)


def duplicate_identity(value: object) -> str:
    return normalize_text(value).casefold()


def workbook_header_maps(ws) -> tuple[dict[str, int], dict[str, int]]:
    raw: dict[str, int] = {}
    canonical: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(1, col_idx).value)
        if not value:
            continue
        raw[value] = col_idx
        key = canonical_key_for_header(value)
        if key:
            canonical[key] = col_idx
    return raw, canonical


def normalize_headers(workbook_path: Path, output_path: Path | None = None) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    output_path = (output_path or workbook_path.with_name(f"{workbook_path.stem}.normalized{workbook_path.suffix}")).resolve()
    if output_path != workbook_path:
        output_path.write_bytes(workbook_path.read_bytes())

    wb = load_workbook(output_path)
    header_changes: list[dict[str, object]] = []
    for ws in wb.worksheets:
        raw_headers = [normalize_text(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
        detected_language = detect_header_language(raw_headers)
        for col_idx in range(1, ws.max_column + 1):
            current = normalize_text(ws.cell(1, col_idx).value)
            if not current:
                continue
            key = canonical_key_for_header(current)
            if key:
                new_value = preferred_header_for_key(key, language="english")
            else:
                fg_key = canonical_fieldguide_key(current)
                new_value = fg_key if fg_key else None
            if new_value and new_value != current:
                ws.cell(1, col_idx).value = new_value
                header_changes.append(
                    {
                        "sheet": ws.title,
                        "column": col_idx,
                        "from": current,
                        "to": new_value,
                        "detected_language": detected_language,
                    }
                )
    wb.save(output_path)
    mapping_path = project_file_paths(output_path)["header_mapping"]
    write_json(mapping_path, {"workbook": str(output_path), "changes": header_changes})
    return {"workbook": str(output_path), "header_changes": header_changes, "mapping_path": str(mapping_path)}


def append_rows(workbook_path: Path, sheet_name: str, rows: list[dict[str, object]]) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    wb = load_workbook(workbook_path)
    ws = wb[sheet_name]
    _, canonical_headers = workbook_header_maps(ws)

    def ensure_header(key: str) -> int:
        col = canonical_headers.get(key)
        if col:
            return col
        col = ws.max_column + 1
        ws.cell(1, col).value = preferred_header_for_key(key, language="english")
        canonical_headers[key] = col
        return col

    title_col = ensure_header("paper_title")
    link_col = ensure_header("paper_link")
    abstract_col = ensure_header("abstract")
    writing_section_col = ensure_header("writing_section")
    writing_argument_col = ensure_header("writing_argument")

    existing_titles: dict[str, int] = {}
    existing_links: dict[str, int] = {}
    for row_idx in range(2, ws.max_row + 1):
        existing_title = duplicate_identity(ws.cell(row_idx, title_col).value)
        existing_link = duplicate_identity(ws.cell(row_idx, link_col).value)
        if existing_title and existing_title not in existing_titles:
            existing_titles[existing_title] = row_idx
        if existing_link and existing_link not in existing_links:
            existing_links[existing_link] = row_idx

    pending_titles: dict[str, int] = {}
    pending_links: dict[str, int] = {}
    inserted_items: list[dict[str, object]] = []
    added_rows: list[dict[str, object]] = []
    skipped_rows: list[dict[str, object]] = []
    for item in rows:
        title = normalize_text(item.get("title"))
        link = normalize_text(item.get("link"))
        if not title:
            skipped_rows.append(
                {
                    "row": None,
                    "title": title,
                    "link": link,
                    "reason": "missing_title",
                }
            )
            continue

        title_id = duplicate_identity(title)
        link_id = duplicate_identity(link)

        if title_id and title_id in existing_titles:
            skipped_rows.append(
                {
                    "row": existing_titles[title_id],
                    "title": title,
                    "link": link,
                    "reason": "duplicate_title",
                }
            )
            continue
        if link_id and link_id in existing_links:
            skipped_rows.append(
                {
                    "row": existing_links[link_id],
                    "title": title,
                    "link": link,
                    "reason": "duplicate_link",
                }
            )
            continue
        if title_id and title_id in pending_titles:
            skipped_rows.append(
                {
                    "row": pending_titles[title_id],
                    "title": title,
                    "link": link,
                    "reason": "duplicate_title_in_payload",
                }
            )
            continue
        if link_id and link_id in pending_links:
            skipped_rows.append(
                {
                    "row": pending_links[link_id],
                    "title": title,
                    "link": link,
                    "reason": "duplicate_link_in_payload",
                }
            )
            continue

        row_idx = ws.max_row + 1
        ws.cell(row_idx, title_col).value = title
        ws.cell(row_idx, link_col).value = link
        ws.cell(row_idx, abstract_col).value = normalize_text(item.get("abstract"))
        ws.cell(row_idx, writing_section_col).value = normalize_text(item.get("writing_section"))
        ws.cell(row_idx, writing_argument_col).value = normalize_text(item.get("writing_argument"))

        if title_id:
            pending_titles[title_id] = row_idx
        if link_id:
            pending_links[link_id] = row_idx
        inserted_items.append({"row": row_idx, "title": title, "link": link})

    wb.save(workbook_path)

    wb_verify = load_workbook(workbook_path, read_only=True)
    ws_verify = wb_verify[sheet_name]
    _, verified_headers = workbook_header_maps(ws_verify)
    verified_title_col = verified_headers.get("paper_title")
    verified_link_col = verified_headers.get("paper_link")

    for item in inserted_items:
        title = item["title"]
        link = item["link"]
        title_id = duplicate_identity(title)
        link_id = duplicate_identity(link)
        verified_row = None
        for row_idx in range(2, ws_verify.max_row + 1):
            current_title = duplicate_identity(ws_verify.cell(row_idx, verified_title_col).value if verified_title_col else "")
            current_link = duplicate_identity(ws_verify.cell(row_idx, verified_link_col).value if verified_link_col else "")
            if title_id and current_title == title_id:
                verified_row = row_idx
                break
            if link_id and current_link == link_id:
                verified_row = row_idx
                break

        added_rows.append({"row": verified_row, "title": title, "link": link})

    return {
        "workbook": str(workbook_path),
        "sheet_name": sheet_name,
        "added_rows": added_rows,
        "skipped_rows": skipped_rows,
    }


def writeback_rows(workbook_path: Path, sheet_name: str | None, preserve_headers: bool) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    wb = load_workbook(workbook_path)
    chosen_sheet = sheet_name or wb.sheetnames[0]
    ws = wb[chosen_sheet]
    raw_headers, canonical_headers = workbook_header_maps(ws)
    header_language = detect_header_language(list(raw_headers.keys()))

    def ensure_header(key: str) -> int:
        col = canonical_headers.get(key)
        if col:
            return col
        col = ws.max_column + 1
        ws.cell(1, col).value = preferred_header_for_key(key, language=header_language if preserve_headers else "english")
        canonical_headers[key] = col
        return col

    updated_rows: list[int] = []
    analysis_dir = workbook_path.parent / f"{workbook_path.stem}_artifacts" / "row-analysis"
    for path in sorted(analysis_dir.glob("*.json")):
        payload = read_json_if_exists(path, {})
        if not isinstance(payload, dict):
            continue
        row = int(payload["row"])
        final_fields = payload.get("final_fields") or {}
        for key, value in final_fields.items():
            ws.cell(row, ensure_header(key)).value = value
        if payload.get("writing_section"):
            ws.cell(row, ensure_header("writing_section")).value = payload["writing_section"]
        if payload.get("writing_argument"):
            ws.cell(row, ensure_header("writing_argument")).value = payload["writing_argument"]
        warnings = payload.get("warnings") or []
        if warnings:
            ws.cell(row, ensure_header("verification_status")).value = " | ".join(str(item) for item in warnings)
        local_path = normalize_text(final_fields.get("local_file_path"))
        if local_path:
            cell = ws.cell(row, ensure_header("local_file_path"))
            rel = relative_path(Path(local_path), workbook_path.parent)
            cell.value = rel
            if Path(local_path).exists():
                cell.hyperlink = rel
                cell.style = "Hyperlink"
        evidence_path = normalize_text(final_fields.get("evidence_path"))
        if evidence_path:
            cell = ws.cell(row, ensure_header("evidence_path"))
            rel = relative_path(Path(evidence_path), workbook_path.parent)
            cell.value = rel
            if Path(evidence_path).exists():
                cell.hyperlink = rel
                cell.style = "Hyperlink"
        updated_rows.append(row)

    wb.save(workbook_path)
    return {
        "workbook": str(workbook_path),
        "sheet_name": chosen_sheet,
        "updated_rows": updated_rows,
        "header_language": header_language,
    }
