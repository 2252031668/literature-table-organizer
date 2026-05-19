from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from common import (
    STATUS_PDF_READY,
    ensure_dir,
    normalize_text,
    paper_reports_dir_for_workbook,
    project_file_paths,
    row_analysis_dir_for_workbook,
    slugify,
)
from lib.project import choose_primary_sheet
from lib.workbook import workbook_header_maps


def build_report_template(row: int, title: str, link: str, pdf_path: Path) -> str:
    return f"""# Paper Report

## Part A - 中文阅读报告

### Basic Information
- Row: {row}
- Title: {title}
- Link: {link}
- PDF Path: {pdf_path}

### 这篇论文解决什么问题
- 待填写

### 方法是什么，亮点是什么
- 待填写

### 实验怎么做
- 待填写

### 结果最关键的点是什么
- 待填写

### 局限和展望
- 待填写

### 这篇论文对本综述有什么意义
- 待填写

## Part B - English Field Decisions

### field: paradigm_mapping
- decision:
- reasoning:
- confidence:
- warning:

### field: paradigm_subtype
- decision:
- reasoning:
- confidence:
- warning:

### field: functional_layer
- decision:
- reasoning:
- confidence:
- warning:

### field: planning_granularity
- decision:
- reasoning:
- confidence:
- warning:

### field: bimanual_design_type
- decision:
- reasoning:
- confidence:
- warning:

### field: verified_core_method
- decision:
- reasoning:
- confidence:
- warning:

### field: verified_tasks_datasets
- decision:
- reasoning:
- confidence:
- warning:

### field: verified_key_results
- decision:
- reasoning:
- confidence:
- warning:

### field: verified_limitations
- decision:
- reasoning:
- confidence:
- warning:

### field: writing_section
- decision:
- reasoning:
- confidence:
- warning:

### field: writing_argument
- decision:
- reasoning:
- confidence:
- warning:
"""


def create_row_analysis_shell(workbook_path: Path, row: int, pdf_path: Path, sheet_name: str | None = None) -> dict[str, object]:
    workbook_path = workbook_path.resolve()
    pdf_path = pdf_path.resolve()
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF does not exist: {pdf_path}")

    framework_path = project_file_paths(workbook_path)["framework"]
    if not framework_path.exists():
        raise FileNotFoundError("framework.md does not exist. Run init-project first.")

    wb = load_workbook(workbook_path, read_only=True)
    chosen_sheet = sheet_name or choose_primary_sheet(workbook_path)
    ws = wb[chosen_sheet]
    _, canonical_headers = workbook_header_maps(ws)

    def cell_for(key: str) -> str:
        col = canonical_headers.get(key)
        if not col:
            return ""
        return normalize_text(ws.cell(row, col).value)

    title = cell_for("paper_title")
    if not title:
        raise ValueError(f"Row {row} has no paper title in sheet {chosen_sheet}")

    row_slug = slugify(title)
    reports_dir = ensure_dir(paper_reports_dir_for_workbook(workbook_path))
    analysis_dir = ensure_dir(row_analysis_dir_for_workbook(workbook_path))
    report_path = reports_dir / f"{row:03d}-{row_slug}.md"
    json_path = analysis_dir / f"{row:03d}-{row_slug}.json"

    report_path.write_text(build_report_template(row, title, cell_for("paper_link"), pdf_path), encoding="utf-8")

    payload = {
        "row": row,
        "title": title,
        "pdf_path": str(pdf_path),
        "resolved_link": cell_for("paper_link"),
        "source_status": STATUS_PDF_READY,
        "review_required": True,
        "warnings": [
            "Read the PDF directly before filling decisions.",
            "Do not rely on abstract-only evidence.",
        ],
        "final_fields": {
            "paper_title": title,
            "paper_link": cell_for("paper_link"),
            "abstract": cell_for("abstract"),
            "local_file_path": str(pdf_path),
            "evidence_path": str(report_path),
        },
        "writing_section": "",
        "writing_argument": "",
        "confidence_map": {},
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    return {
        "sheet_name": chosen_sheet,
        "row": row,
        "framework_path": str(framework_path),
        "report_path": str(report_path),
        "json_path": str(json_path),
    }
